"""Strażnik bomby ZIP przed openpyxl (audyt 10.10.2026, S14) – sam moduł i import listy uczniów.

Archiwa budujemy w pamięci: ``zipfile`` ściska trzydzieści megabajtów zer do kilkudziesięciu
kilobajtów, więc „bomba” mieści się w limicie wgrania (2 MB) dokładnie tak, jak w ataku – a test
nie musi tworzyć gigabajtów.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook

from apps.core.xlsx import DEFAULT_MAX_UNCOMPRESSED, XlsxBombError, check_xlsx_bomb


def zip_with(name: str, payload: bytes) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(name, payload)
    return buffer.getvalue()


def real_workbook(rows: int = 3) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["imię", "nazwisko", "e-mail", "rok urodzenia", "klasa"])
    for index in range(rows):
        sheet.append(["Kasia", "Nowak", f"k{index}@example.test", 2000, 2])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_zwykly_arkusz_przechodzi():
    check_xlsx_bomb(real_workbook(rows=500))


def test_archiwum_ponad_limit_po_rozpakowaniu_jest_odmowa():
    bomb = zip_with("xl/sharedStrings.xml", b"\0" * (DEFAULT_MAX_UNCOMPRESSED + 10 * 1024 * 1024))
    # Warunek ataku: plik mieści się w limicie wgrania, a po rozpakowaniu go przekracza.
    assert len(bomb) < 2 * 1024 * 1024
    with pytest.raises(XlsxBombError):
        check_xlsx_bomb(bomb)


def test_podejrzanie_wysoki_stopien_kompresji_jest_odmowa():
    """Pięć megabajtów zer mieści się w limicie sumy, ale ściska się ~1000 razy – plik spreparowany."""
    with pytest.raises(XlsxBombError):
        check_xlsx_bomb(zip_with("xl/worksheets/sheet1.xml", b"\0" * (5 * 1024 * 1024)))


def test_maly_mocno_scisniety_plik_nie_jest_falszywym_alarmem():
    check_xlsx_bomb(zip_with("xl/styles.xml", b"\0" * (512 * 1024)))


def test_nie_archiwum_jest_odmowa():
    with pytest.raises(XlsxBombError):
        check_xlsx_bomb(b"to nie jest zip")


def test_sciezka_tez_jest_sprawdzana(tmp_path):
    path = tmp_path / "wykaz.xlsx"
    path.write_bytes(zip_with("xl/sharedStrings.xml", b"\0" * (30 * 1024 * 1024)))
    with pytest.raises(XlsxBombError):
        check_xlsx_bomb(path)
    # Hojniejszy limit (wykaz SIO) nadal łapie stopień kompresji.
    with pytest.raises(XlsxBombError):
        check_xlsx_bomb(path, max_uncompressed=1024 * 1024 * 1024)


# --- import listy uczniów -------------------------------------------------------------------------


class Upload(io.BytesIO):
    def __init__(self, data: bytes, name: str):
        super().__init__(data)
        self.name = name
        self.size = len(data)


def test_import_uczniow_odrzuca_bombe_zanim_otworzy_ja_openpyxl():
    from apps.accounts.bulk_registration import read_table
    from apps.core.api import DomainError

    bomb = zip_with("xl/sharedStrings.xml", b"\0" * (30 * 1024 * 1024))
    with pytest.raises(DomainError) as exc:
        read_table(Upload(bomb, "lista.xlsx"))
    assert exc.value.machine_code == "IMPORT_XLSX_TOO_LARGE"


def test_import_uczniow_przestaje_czytac_po_limicie_wierszy():
    from apps.accounts.bulk_registration import MAX_ROWS, _xlsx_table

    table = _xlsx_table(real_workbook(rows=MAX_ROWS + 300))
    # Nagłówek + limit + jeden ponad: tyle wystarczy, żeby ``parse_table`` odmówił pliku.
    assert len(table) == MAX_ROWS + 2


def test_import_uczniow_nie_trzyma_pustych_wierszy_z_konca_arkusza():
    from apps.accounts.bulk_registration import _xlsx_table

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["imię", "nazwisko"])
    sheet.append([])
    sheet.append(["Kasia", "Nowak"])
    sheet.cell(row=5000, column=1, value=None)
    sheet.cell(row=5000, column=2, value="")
    buffer = io.BytesIO()
    workbook.save(buffer)
    table = _xlsx_table(buffer.getvalue())
    # Pusty wiersz w środku zostaje (numer wiersza w podglądzie), puste na końcu – nie.
    assert [row[:2] for row in table] == [["imię", "nazwisko"], [], ["Kasia", "Nowak"]]
