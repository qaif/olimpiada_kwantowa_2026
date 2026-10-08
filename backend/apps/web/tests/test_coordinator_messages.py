"""Ekran ``/coordinator/messages/`` – wysyłka komunikatów organizatora.

Reguły doboru odbiorców i podziału na porcje mają własne testy w
``apps/accounts/tests/test_messaging.py``. Tutaj sprawdzamy sam ekran, a w nim przede wszystkim
jedyny bezpiecznik przed omyłkową wysyłką do kilku tysięcy osób: **podgląd nic nie wysyła**.
"""

import re

import pytest
from django.core import mail

from apps.accounts.models import BroadcastGroup, BroadcastStatus, MessageBroadcast, Region
from apps.accounts.tests.factories import ParticipantFactory, UserFactory
from apps.competitions.tests.factories import StageEntryFactory
from apps.core.models import AuditLog

pytestmark = pytest.mark.django_db

BASE = {"subject": "Zmiana terminu", "body": "Etap rusza tydzień później."}


SIGNATURE = re.compile(r'name="preview_signature" value="([0-9a-f]+)"')


def _post(client, data, *, action="preview"):
    """Podgląd albo – dla ``action="send"`` – dokładnie to, co robi koordynator: podgląd, potem „Wyślij”.

    Wysyłka bez podglądu nie przechodzi (podpis podglądu), więc pomocnik najpierw ogląda, zabiera
    podpis z ukrytego pola i dopiero z nim wysyła. Przypadki „bez podglądu” i „coś zmieniono po
    podglądzie” mają własne testy niżej.
    """
    if action != "send":
        return client.post("/coordinator/messages/", {**data, "action": action})
    preview = client.post("/coordinator/messages/", {**data, "action": "preview"}).content.decode()
    match = SIGNATURE.search(preview)
    signature = match.group(1) if match else ""
    return client.post("/coordinator/messages/", {**data, "action": "send", "preview_signature": signature})


def test_preview_shows_the_count_and_sends_nothing(web_client, coordinator, elim_stage, entry):
    web_client.force_login(coordinator)
    mail.outbox.clear()

    response = _post(web_client, {**BASE, "group": BroadcastGroup.EDITION_PARTICIPANTS})
    content = response.content.decode()

    assert response.status_code == 200
    assert "1 odbiorców" in content
    # Treść widać dokładnie tak, jak pójdzie w liście – to jedyny moment, w którym pomyłkę
    # da się jeszcze cofnąć.
    assert BASE["body"] in content
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def test_send_queues_one_letter_per_recipient_and_records_the_broadcast(
    web_client, coordinator, elim_stage, entry, django_capture_on_commit_callbacks
):
    second = ParticipantFactory(user=UserFactory(email="druga@example.test"))
    StageEntryFactory(participant=second, stage=elim_stage)
    web_client.force_login(coordinator)
    mail.outbox.clear()

    with django_capture_on_commit_callbacks(execute=True):
        response = _post(web_client, {**BASE, "group": BroadcastGroup.EDITION_PARTICIPANTS}, action="send")

    assert response.status_code == 302
    # Jeden list na odbiorcę – nigdy jedna koperta z listą adresów w nagłówku.
    assert len(mail.outbox) == 2
    assert all(len(message.to) == 1 for message in mail.outbox)
    broadcast = MessageBroadcast.objects.get()
    assert broadcast.recipient_count == 2
    assert broadcast.status == BroadcastStatus.SENT
    assert broadcast.created_by == coordinator


def test_audit_records_counters_but_no_addresses(
    web_client, coordinator, elim_stage, entry, django_capture_on_commit_callbacks
):
    web_client.force_login(coordinator)

    with django_capture_on_commit_callbacks(execute=True):
        _post(web_client, {**BASE, "group": BroadcastGroup.EDITION_PARTICIPANTS}, action="send")

    log = AuditLog.objects.get(action="broadcast.sent")
    assert log.diff["recipients"] == 1
    assert entry.participant.user.email not in str(log.diff)


def test_custom_list_requires_at_least_one_address(web_client, coordinator):
    web_client.force_login(coordinator)

    response = _post(web_client, {**BASE, "group": BroadcastGroup.CUSTOM, "addresses": "  "})

    assert response.status_code == 200
    assert "Wklej przynajmniej jeden adres." in response.content.decode()


def test_stage_group_requires_a_stage(web_client, coordinator, elim_stage):
    web_client.force_login(coordinator)

    response = _post(web_client, {**BASE, "group": BroadcastGroup.STAGE_REGISTERED})

    assert response.status_code == 200
    assert "wymaga wskazania etapu" in response.content.decode()


def test_sending_to_an_empty_group_records_nothing(
    web_client, coordinator, elim_stage, django_capture_on_commit_callbacks
):
    web_client.force_login(coordinator)
    mail.outbox.clear()

    with django_capture_on_commit_callbacks(execute=True):
        response = _post(
            web_client,
            {**BASE, "group": BroadcastGroup.STAGE_REGISTERED, "stage": elim_stage.pk},
            action="send",
        )

    assert response.status_code == 200
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def test_history_lists_past_broadcasts(
    web_client, coordinator, elim_stage, entry, django_capture_on_commit_callbacks
):
    web_client.force_login(coordinator)
    with django_capture_on_commit_callbacks(execute=True):
        _post(web_client, {**BASE, "group": BroadcastGroup.EDITION_PARTICIPANTS}, action="send")

    response = web_client.get("/coordinator/messages/")

    assert BASE["subject"] in response.content.decode()


def test_send_without_a_preview_sends_nothing_and_shows_the_preview(
    web_client, coordinator, elim_stage, entry, django_capture_on_commit_callbacks
):
    web_client.force_login(coordinator)
    mail.outbox.clear()

    with django_capture_on_commit_callbacks(execute=True):
        response = web_client.post(
            "/coordinator/messages/", {**BASE, "group": BroadcastGroup.EDITION_PARTICIPANTS, "action": "send"}
        )

    content = response.content.decode()
    assert response.status_code == 200
    assert "zmieniły się od podglądu" in content
    assert "1 odbiorców" in content
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def test_switching_the_group_after_the_preview_does_not_send(
    web_client, coordinator, elim_stage, entry, django_capture_on_commit_callbacks
):
    """Podgląd listu do jednego etapu nie jest przepustką dla listu do wszystkich uczestników."""
    ParticipantFactory(user=UserFactory(email="bez-wpisu@example.test"))
    web_client.force_login(coordinator)
    mail.outbox.clear()
    narrow = {**BASE, "group": BroadcastGroup.STAGE_REGISTERED, "stage": elim_stage.pk}
    signature = SIGNATURE.search(_post(web_client, narrow).content.decode()).group(1)

    with django_capture_on_commit_callbacks(execute=True):
        response = web_client.post(
            "/coordinator/messages/",
            {
                **narrow,
                "group": BroadcastGroup.ALL_PARTICIPANTS,
                "action": "send",
                "preview_signature": signature,
            },
        )

    assert response.status_code == 200
    assert "2 odbiorców" in response.content.decode()
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def test_page_is_forbidden_for_a_reviewer(web_client, reviewer):
    web_client.force_login(reviewer.user)

    assert web_client.get("/coordinator/messages/").status_code == 403


def test_dashboard_links_to_the_messages_screen(web_client, coordinator):
    web_client.force_login(coordinator)

    assert "/coordinator/messages/" in web_client.get("/coordinator/").content.decode()


# --- grupy z 24.09.2026 ---------------------------------------------------------------------------


def _options(content: str, name: str) -> list[str]:
    """Wartości ``<option>`` listy wyboru ``name`` – w kolejności, w jakiej widzi je koordynator."""
    select = re.search(rf'<select name="{name}"[^>]*>(.*?)</select>', content, re.S)
    assert select, f"brak listy {name}"
    return re.findall(r'<option value="([^"]*)"', select.group(1))


def test_all_participants_is_the_first_group_on_the_list(web_client, coordinator):
    web_client.force_login(coordinator)

    content = web_client.get("/coordinator/messages/").content.decode()

    assert _options(content, "group")[0] == BroadcastGroup.ALL_PARTICIPANTS


def test_all_participants_preview_counts_people_without_stage_entries_and_nobody_else(
    web_client, coordinator, entry, other_competition
):
    ParticipantFactory(user=UserFactory(email="bez-wpisu@example.test"))
    ParticipantFactory(user=UserFactory(email="sasiad@example.test"), competition=other_competition)
    web_client.force_login(coordinator)
    mail.outbox.clear()

    content = _post(web_client, {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS}).content.decode()

    assert "2 odbiorców" in content
    assert "wszyscy uczestnicy konkursu" in content
    assert not mail.outbox


@pytest.mark.parametrize(
    ("group", "message"),
    [
        (BroadcastGroup.STAGE_NO_SUBMISSION, "wymaga wskazania etapu"),
        (BroadcastGroup.REGION_PARTICIPANTS, "Wybierz województwo uczestników."),
        (BroadcastGroup.COMMITTEE_DISTRICT, "Wybierz województwo komitetu."),
        (BroadcastGroup.SCHOOL_PARTICIPANTS, "Wybierz szkołę."),
        (BroadcastGroup.GRADE_PARTICIPANTS, "Wybierz klasę."),
        (BroadcastGroup.WORKSHOP_ATTENDEES, "Wybierz warsztat."),
    ],
)
def test_group_with_a_parameter_requires_it(web_client, coordinator, group, message):
    web_client.force_login(coordinator)
    mail.outbox.clear()

    response = _post(web_client, {**BASE, "group": group}, action="send")

    assert response.status_code == 200
    assert message in response.content.decode()
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def test_school_of_another_competition_is_not_a_valid_choice(web_client, coordinator, other_competition):
    from apps.schools.tests.factories import SchoolFactory

    school = SchoolFactory()
    ParticipantFactory(competition=other_competition, school_ref=school)
    web_client.force_login(coordinator)

    response = _post(
        web_client, {**BASE, "group": BroadcastGroup.SCHOOL_PARTICIPANTS, "school": f"sio:{school.pk}"}
    )

    assert "Wybierz poprawną wartość" in response.content.decode()


def test_region_screen_uses_custom_regions_when_the_flag_is_on(web_client, coordinator, competition, edition):
    north = Region.objects.create(competition=competition, code="okreg-polnoc", name="Okręg Północ")
    ParticipantFactory(user=UserFactory(email="polnoc@example.test"), region=north)
    web_client.force_login(coordinator)
    assert 'name="region"' not in web_client.get("/coordinator/messages/").content.decode()

    competition.feature_flags = {**(competition.feature_flags or {}), "custom_regions": True}
    competition.save(update_fields=["feature_flags"])
    missing = _post(web_client, {**BASE, "group": BroadcastGroup.REGION_PARTICIPANTS}).content.decode()
    preview = _post(
        web_client, {**BASE, "group": BroadcastGroup.REGION_PARTICIPANTS, "region": north.pk}
    ).content.decode()

    assert "Wybierz region uczestników." in missing
    assert "1 odbiorców" in preview
    assert "Okręg Północ" in preview


def test_send_to_a_school_records_which_school_in_the_history(
    web_client, coordinator, edition, django_capture_on_commit_callbacks
):
    from apps.schools.tests.factories import SchoolFactory

    school = SchoolFactory(name="XIV LO im. Staszica", city="Warszawa")
    ParticipantFactory(user=UserFactory(email="staszic@example.test"), school_ref=school)
    ParticipantFactory(user=UserFactory(email="inna@example.test"))
    web_client.force_login(coordinator)
    mail.outbox.clear()
    data = {**BASE, "group": BroadcastGroup.SCHOOL_PARTICIPANTS, "school": f"sio:{school.pk}"}

    preview = _post(web_client, data).content.decode()
    with django_capture_on_commit_callbacks(execute=True):
        _post(web_client, data, action="send")

    assert "XIV LO im. Staszica, Warszawa" in preview
    assert [message.to for message in mail.outbox] == [["staszic@example.test"]]
    broadcast = MessageBroadcast.objects.get()
    label = "XIV LO im. Staszica, Warszawa (bieżąca edycja)"
    assert broadcast.target == {"school": f"sio:{school.pk}", "past_editions": False, "label": label}
    history = web_client.get("/coordinator/messages/").content.decode()
    assert label in history
    log = AuditLog.objects.get(action="broadcast.sent")
    assert log.diff["target"]["label"] == label


def test_stage_parameter_is_recorded_for_the_no_submission_reminder(
    web_client, coordinator, elim_stage, entry, django_capture_on_commit_callbacks
):
    web_client.force_login(coordinator)
    data = {**BASE, "group": BroadcastGroup.STAGE_NO_SUBMISSION, "stage": elim_stage.pk}

    with django_capture_on_commit_callbacks(execute=True):
        _post(web_client, data, action="send")

    broadcast = MessageBroadcast.objects.get()
    assert broadcast.recipient_count == 1
    assert broadcast.target == {"stage": elim_stage.pk, "label": str(elim_stage)}


def test_parameter_of_another_group_is_ignored(
    web_client, coordinator, elim_stage, entry, django_capture_on_commit_callbacks
):
    """Etap wybrany „przy okazji” nie zawęża listu do wszystkich i nie trafia do historii."""
    ParticipantFactory(user=UserFactory(email="bez-wpisu@example.test"))
    web_client.force_login(coordinator)

    with django_capture_on_commit_callbacks(execute=True):
        _post(
            web_client,
            {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS, "stage": elim_stage.pk},
            action="send",
        )

    broadcast = MessageBroadcast.objects.get()
    assert broadcast.recipient_count == 2
    # Z parametrów zostaje wyłącznie zakres edycji – jedyne pole, które należy do tej grupy.
    assert broadcast.target == {"past_editions": False, "label": "bieżąca edycja"}


def test_workshop_group_on_the_screen(web_client, coordinator, django_capture_on_commit_callbacks):
    from datetime import date

    from apps.cms.models import ContentPage, HomePage, WorkshopAttendance
    from apps.cms.workshops import WORKSHOPS_SLUG, workshop_rows

    home = HomePage.objects.get()
    page = ContentPage(
        title="Warsztaty",
        slug=WORKSHOPS_SLUG,
        live=True,
        body=[
            (
                "schedule",
                {
                    "caption": "",
                    "topic_label": "Temat",
                    "date_label": "Termin",
                    "time_label": "",
                    "lecturer_label": "Prowadzący",
                    "rows": [
                        {
                            "topic": "Kubity i bramki",
                            "date": "12.11.2026",
                            "date_value": date(2026, 11, 12),
                            "time": "",
                            "lecturer": "",
                        }
                    ],
                },
            )
        ],
    )
    home.add_child(instance=page)
    page.save_revision().publish()
    key = workshop_rows(page)[0]["key"]
    present = ParticipantFactory(user=UserFactory(email="obecna@example.test"))
    ParticipantFactory(user=UserFactory(email="nieobecna@example.test"))
    WorkshopAttendance.objects.create(participant=present, workshop_key=key)
    web_client.force_login(coordinator)

    assert key in _options(web_client.get("/coordinator/messages/").content.decode(), "workshop")
    with django_capture_on_commit_callbacks(execute=True):
        _post(
            web_client, {**BASE, "group": BroadcastGroup.WORKSHOP_ATTENDEES, "workshop": key}, action="send"
        )

    broadcast = MessageBroadcast.objects.get()
    assert broadcast.recipient_count == 1
    assert broadcast.target == {"workshop": key, "label": "12.11.2026 – Kubity i bramki"}


def test_history_hides_broadcasts_of_another_competition(
    web_client, coordinator, competition, other_competition
):
    from apps.tenancy.tests.factories import create_scoped

    create_scoped(
        MessageBroadcast, competition, group=BroadcastGroup.CUSTOM, subject="Nasz komunikat", body="."
    )
    create_scoped(
        MessageBroadcast, other_competition, group=BroadcastGroup.CUSTOM, subject="Cudzy komunikat", body="."
    )
    web_client.force_login(coordinator)

    content = web_client.get("/coordinator/messages/").content.decode()

    assert "Nasz komunikat" in content
    assert "Cudzy komunikat" not in content


def test_broadcast_belongs_to_the_competition_of_the_request(
    client_for, other_competition, django_capture_on_commit_callbacks
):
    """Koordynator konkursu B wysyła z domeny B – wiersz rejestru należy do B, odbiorcy też są z B."""
    from apps.accounts.models import GROUP_COORDINATOR
    from apps.accounts.tests.factories import CoordinatorFactory
    from apps.tenancy.tests.factories import grant_membership

    coordinator = CoordinatorFactory()
    grant_membership(coordinator, other_competition, GROUP_COORDINATOR)
    ParticipantFactory(user=UserFactory(email="tutejszy@example.test"))
    ParticipantFactory(user=UserFactory(email="w-b@example.test"), competition=other_competition)
    client = client_for(other_competition)
    client.force_login(coordinator)
    mail.outbox.clear()

    with django_capture_on_commit_callbacks(execute=True):
        # Konkurs B nie ma bieżącej edycji – bez przełącznika grupa byłaby pusta.
        response = _post(
            client,
            {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS, "include_past_editions": "on"},
            action="send",
        )

    assert response.status_code == 302
    assert [message.to for message in mail.outbox] == [["w-b@example.test"]]
    assert MessageBroadcast.objects.get().competition == other_competition


# --- zakres edycji (decyzja organizatora z 24.09.2026) ----------------------------------------------


def _veteran(email: str):
    """Uczestnik z konta sprzed bieżącej edycji, bez wpisu do jej etapów."""
    from datetime import timedelta

    from django.utils import timezone

    return ParticipantFactory(user=UserFactory(email=email, date_joined=timezone.now() - timedelta(days=400)))


def test_all_participants_default_to_the_current_edition_and_the_switch_widens_it(
    web_client, coordinator, entry, django_capture_on_commit_callbacks
):
    _veteran("zeszloroczny@example.test")
    web_client.force_login(coordinator)
    mail.outbox.clear()
    data = {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS}

    narrow = _post(web_client, data).content.decode()
    wide = _post(web_client, {**data, "include_past_editions": "on"}).content.decode()
    with django_capture_on_commit_callbacks(execute=True):
        _post(web_client, {**data, "include_past_editions": "on"}, action="send")

    assert "1 odbiorców" in narrow
    assert "wszyscy uczestnicy konkursu: bieżąca edycja" in narrow
    assert "2 odbiorców" in wide
    assert "wszyscy uczestnicy konkursu: także poprzednie edycje" in wide
    broadcast = MessageBroadcast.objects.get()
    assert broadcast.recipient_count == 2
    assert broadcast.target == {"past_editions": True, "label": "także poprzednie edycje"}
    assert AuditLog.objects.get(action="broadcast.sent").diff["target"]["past_editions"] is True
    assert "także poprzednie edycje" in web_client.get("/coordinator/messages/").content.decode()


def test_ticking_past_editions_after_the_preview_does_not_send(
    web_client, coordinator, entry, django_capture_on_commit_callbacks
):
    """Podgląd „bieżąca edycja” nie jest przepustką dla listu do wszystkich roczników."""
    _veteran("zeszloroczny@example.test")
    web_client.force_login(coordinator)
    mail.outbox.clear()
    data = {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS}
    signature = SIGNATURE.search(_post(web_client, data).content.decode()).group(1)

    with django_capture_on_commit_callbacks(execute=True):
        response = web_client.post(
            "/coordinator/messages/",
            {**data, "include_past_editions": "on", "action": "send", "preview_signature": signature},
        )

    assert "zmieniły się od podglądu" in response.content.decode()
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def test_past_editions_switch_is_offered_only_to_the_groups_it_affects(web_client, coordinator):
    import json

    web_client.force_login(coordinator)
    content = web_client.get("/coordinator/messages/").content.decode()
    mapping = json.loads(re.search(r'id="broadcast-parameters"[^>]*>(.*?)</script>', content, re.S).group(1))

    with_switch = {group for group, fields in mapping.items() if "include_past_editions" in fields}
    assert with_switch == {
        BroadcastGroup.ALL_PARTICIPANTS,
        BroadcastGroup.ALL_PARTICIPANTS_AND_TEACHERS,
        BroadcastGroup.REGION_PARTICIPANTS,
        BroadcastGroup.SCHOOL_PARTICIPANTS,
        BroadcastGroup.GRADE_PARTICIPANTS,
    }
    assert 'data-broadcast-param="include_past_editions"' in content


# --- „Eksportuj do Excela” (MSG-EXPORT-01) ---------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _export(client, data):
    """Kliknięcie „Eksportuj do Excela”: ten sam formularz, ``action=export``, bez tematu i treści."""
    return client.post("/coordinator/messages/", {**data, "action": "export"})


def _sheet(response) -> list[tuple]:
    """Wiersze pierwszego arkusza pobranego pliku (z nagłówkiem)."""
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    return [tuple(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]


def _teacher(email: str, competition, **user_kwargs):
    from apps.accounts.models import GROUP_SUPERVISOR, SchoolSupervisor

    user = UserFactory(email=email, groups=[GROUP_SUPERVISOR], **user_kwargs)
    return SchoolSupervisor.objects.create(user=user, school="XIV LO", competition=competition)


def test_messages_page_offers_the_export_button(web_client, coordinator):
    web_client.force_login(coordinator)

    content = web_client.get("/coordinator/messages/").content.decode()

    assert "Eksportuj do Excela" in content
    # Temat i treść nie są do pliku potrzebne – przeglądarka nie może zatrzymać przycisku na nich.
    assert re.search(r'value="export"[^>]*formnovalidate', content)
    assert f'data-broadcast-export="{BroadcastGroup.CUSTOM}"' in content


def test_export_of_participants_and_teachers_downloads_one_named_row_per_person(
    web_client, coordinator, competition, other_competition
):
    from apps.accounts.messaging import resolve_recipients
    from apps.competitions.services import current_edition

    ParticipantFactory(user=UserFactory(email="uczen@example.test", first_name="Ala", last_name="Zielińska"))
    _teacher("nauczyciel@example.test", competition, first_name="Ewa", last_name="Bąk")
    # Ta sama osoba jako uczestnik i nauczyciel – jeden wiersz.
    both = _teacher("oba@example.test", competition, first_name="Jan", last_name="Mazur")
    ParticipantFactory(user=both.user)
    ParticipantFactory(user=UserFactory(email="sasiad@example.test"), competition=other_competition)
    _teacher("nauczyciel-sasiada@example.test", other_competition)
    web_client.force_login(coordinator)
    mail.outbox.clear()
    data = {"group": BroadcastGroup.ALL_PARTICIPANTS_AND_TEACHERS, "include_past_editions": "on"}

    response = _export(web_client, data)

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX
    assert 'filename="odbiorcy-wszyscy-uczestnicy-i-nauczyciele-' in response["Content-Disposition"]
    rows = _sheet(response)
    assert rows == [
        ("Imię", "Nazwisko", "E-mail"),
        ("Ewa", "Bąk", "nauczyciel@example.test"),
        ("Jan", "Mazur", "oba@example.test"),
        ("Ala", "Zielińska", "uczen@example.test"),
    ]
    # Te same adresy, do których poszedłby list – i nic nie wysłano.
    assert sorted(row[2] for row in rows[1:]) == resolve_recipients(
        BroadcastGroup.ALL_PARTICIPANTS_AND_TEACHERS,
        competition=competition,
        edition=current_edition(competition),
        include_past_editions=True,
    )
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def test_export_respects_the_group_parameter(web_client, coordinator, elim_stage, entry):
    from apps.competitions.models import StageKind
    from apps.competitions.tests.factories import StageFactory

    other_stage = StageFactory(edition=elim_stage.edition, kind=StageKind.DISTRICT)
    other = ParticipantFactory(user=UserFactory(email="inny-etap@example.test"))
    StageEntryFactory(participant=other, stage=other_stage)
    web_client.force_login(coordinator)

    response = _export(web_client, {"group": BroadcastGroup.STAGE_REGISTERED, "stage": elim_stage.pk})

    assert [row[2] for row in _sheet(response)[1:]] == [entry.participant.user.email]


def test_export_neutralises_formulas_in_names(web_client, coordinator):
    ParticipantFactory(user=UserFactory(email="formula@example.test", first_name='=HYPERLINK("http://x")'))
    web_client.force_login(coordinator)

    rows = _sheet(
        _export(web_client, {"group": BroadcastGroup.ALL_PARTICIPANTS, "include_past_editions": "on"})
    )

    assert rows[1][0] == '\'=HYPERLINK("http://x")'


def test_export_is_audited_without_personal_data(web_client, coordinator, competition, elim_stage, entry):
    web_client.force_login(coordinator)

    _export(web_client, {"group": BroadcastGroup.STAGE_REGISTERED, "stage": elim_stage.pk})

    log = AuditLog.objects.get(action="export.generated")
    assert log.diff["kind"] == "broadcast_recipients"
    assert log.diff["format"] == "xlsx"
    assert log.diff["group"] == BroadcastGroup.STAGE_REGISTERED
    assert log.diff["target"]["stage"] == elim_stage.pk
    assert log.diff["rows"] == 1
    assert log.target_type == "tenancy.competition"
    assert log.target_id == str(competition.pk)
    assert entry.participant.user.email not in str(log.diff)


def test_export_without_the_group_parameter_shows_the_form_errors(web_client, coordinator):
    web_client.force_login(coordinator)

    response = _export(web_client, {"group": BroadcastGroup.SCHOOL_PARTICIPANTS})

    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/html")
    assert "Wybierz szkołę." in response.content.decode()
    assert not AuditLog.objects.filter(action="export.generated").exists()


def test_export_of_a_pasted_list_is_refused(web_client, coordinator):
    web_client.force_login(coordinator)

    response = _export(web_client, {"group": BroadcastGroup.CUSTOM, "addresses": "ktos@example.test"})

    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/html")
    assert "Wklejonej listy adresów nie da się wyeksportować" in response.content.decode()
    assert not AuditLog.objects.filter(action="export.generated").exists()


def test_export_of_an_empty_group_gives_a_message_not_a_file(web_client, coordinator, elim_stage):
    web_client.force_login(coordinator)

    response = _export(web_client, {"group": BroadcastGroup.STAGE_REGISTERED, "stage": elim_stage.pk})

    assert response["Content-Type"].startswith("text/html")
    assert "nie ma czego eksportować" in response.content.decode()
    assert not AuditLog.objects.filter(action="export.generated").exists()


def test_export_is_forbidden_for_a_reviewer(web_client, reviewer, entry):
    web_client.force_login(reviewer.user)

    response = _export(web_client, {"group": BroadcastGroup.ALL_PARTICIPANTS})

    assert response.status_code == 403
    assert not AuditLog.objects.filter(action="export.generated").exists()


def test_export_never_reaches_another_competition(client_for, other_competition):
    """Koordynator konkursu B eksportuje z domeny B – w pliku są wyłącznie ludzie B."""
    from apps.accounts.models import GROUP_COORDINATOR
    from apps.accounts.tests.factories import CoordinatorFactory
    from apps.tenancy.tests.factories import grant_membership

    coordinator = CoordinatorFactory()
    grant_membership(coordinator, other_competition, GROUP_COORDINATOR)
    ParticipantFactory(user=UserFactory(email="tutejszy@example.test"))
    ParticipantFactory(user=UserFactory(email="w-b@example.test"), competition=other_competition)
    client = client_for(other_competition)
    client.force_login(coordinator)

    response = _export(client, {"group": BroadcastGroup.ALL_PARTICIPANTS, "include_past_editions": "on"})

    assert [row[2] for row in _sheet(response)[1:]] == ["w-b@example.test"]
    assert AuditLog.objects.get(action="export.generated").competition == other_competition


# --- komunikaty z datą przyszłą (MSG-SCHED-01) ------------------------------------------------------


def _local(**delta) -> str:
    """Wartość pola ``datetime-local`` (czas polski) przesunięta względem teraz."""
    from datetime import timedelta

    from django.utils import timezone

    return (timezone.localtime() + timedelta(**delta)).strftime("%Y-%m-%dT%H:%M")


def test_preview_with_a_send_time_offers_scheduling_and_warns_about_the_recount(
    web_client, coordinator, entry
):
    web_client.force_login(coordinator)

    content = _post(web_client, {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS, "send_at": _local(days=1)})
    content = content.content.decode()

    assert "1 odbiorców" in content
    assert "Zaplanuj" in content
    assert "policzona ponownie w chwili wysyłki" in content
    assert not MessageBroadcast.objects.exists()


def test_scheduling_from_the_screen_sends_nothing_now(web_client, coordinator, entry):
    web_client.force_login(coordinator)
    mail.outbox.clear()

    response = _post(
        web_client,
        {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS, "send_at": _local(days=2)},
        action="send",
    )

    assert response.status_code == 302
    assert not mail.outbox
    broadcast = MessageBroadcast.objects.get()
    assert broadcast.status == BroadcastStatus.SCHEDULED
    assert broadcast.created_by == coordinator
    assert broadcast.parameters == {"include_past_editions": False}
    assert broadcast.target == {"past_editions": False, "label": "bieżąca edycja"}
    assert AuditLog.objects.get(action="broadcast.scheduled").actor == coordinator
    page = web_client.get("/coordinator/messages/").content.decode()
    assert "Zaplanowane" in page
    assert f"/coordinator/messages/{broadcast.pk}/cancel/" in page


def test_empty_group_can_be_scheduled(web_client, coordinator, elim_stage):
    """Do terminu grupa może się zapełnić – odbiorców i tak liczy beat w chwili wysyłki."""
    web_client.force_login(coordinator)

    response = _post(
        web_client,
        {**BASE, "group": BroadcastGroup.STAGE_REGISTERED, "stage": elim_stage.pk, "send_at": _local(days=1)},
        action="send",
    )

    assert response.status_code == 302
    assert MessageBroadcast.objects.get().parameters == {"stage": elim_stage.pk}


@pytest.mark.parametrize(
    ("delta", "message"),
    [
        ({"hours": -1}, "co najmniej 5 minut"),
        ({"minutes": 2}, "co najmniej 5 minut"),
        ({"days": 91}, "90 dni"),
    ],
)
def test_send_time_too_close_or_too_far_is_refused(web_client, coordinator, entry, delta, message):
    web_client.force_login(coordinator)

    response = _post(
        web_client,
        {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS, "send_at": _local(**delta)},
        action="send",
    )

    assert response.status_code == 200
    assert message in response.content.decode()
    assert not MessageBroadcast.objects.exists()


def test_pasted_list_cannot_be_scheduled_on_the_screen(web_client, coordinator):
    web_client.force_login(coordinator)

    response = _post(
        web_client,
        {**BASE, "group": BroadcastGroup.CUSTOM, "addresses": "a@example.test", "send_at": _local(days=1)},
        action="send",
    )

    assert "Wklejonej listy adresów nie da się zaplanować" in response.content.decode()
    assert not MessageBroadcast.objects.exists()


def test_setting_a_send_time_after_the_preview_does_not_schedule(web_client, coordinator, entry):
    """Podgląd „wyślij od razu” nie jest przepustką dla komunikatu zaplanowanego – ani odwrotnie."""
    web_client.force_login(coordinator)
    mail.outbox.clear()
    data = {**BASE, "group": BroadcastGroup.ALL_PARTICIPANTS}
    signature = SIGNATURE.search(_post(web_client, data).content.decode()).group(1)

    response = web_client.post(
        "/coordinator/messages/",
        {**data, "send_at": _local(days=1), "action": "send", "preview_signature": signature},
    )

    assert "zmieniły się od podglądu" in response.content.decode()
    assert not mail.outbox
    assert not MessageBroadcast.objects.exists()


def _waiting(competition, **kwargs):
    from datetime import timedelta

    from django.utils import timezone

    from apps.tenancy.tests.factories import create_scoped

    return create_scoped(
        MessageBroadcast,
        competition,
        group=BroadcastGroup.COMMITTEE,
        subject=kwargs.pop("subject", "Zaplanowany komunikat"),
        body=".",
        status=BroadcastStatus.SCHEDULED,
        scheduled_for=timezone.now() + timedelta(days=1),
        **kwargs,
    )


def test_cancel_from_the_screen(web_client, coordinator, competition):
    broadcast = _waiting(competition, created_by=coordinator)
    web_client.force_login(coordinator)

    response = web_client.post(f"/coordinator/messages/{broadcast.pk}/cancel/")

    assert response.status_code == 302
    broadcast.refresh_from_db()
    assert broadcast.status == BroadcastStatus.CANCELLED
    assert AuditLog.objects.get(action="broadcast.cancelled").actor == coordinator
    page = web_client.get("/coordinator/messages/").content.decode()
    # Anulowany wypada z „Zaplanowanych” i zostaje w historii ze swoim stanem.
    assert f"/coordinator/messages/{broadcast.pk}/cancel/" not in page
    assert "anulowana" in page


def test_cancel_of_another_competition_is_404(web_client, coordinator, other_competition):
    broadcast = _waiting(other_competition)
    web_client.force_login(coordinator)

    response = web_client.post(f"/coordinator/messages/{broadcast.pk}/cancel/")

    assert response.status_code == 404
    broadcast.refresh_from_db()
    assert broadcast.status == BroadcastStatus.SCHEDULED


def test_cancel_is_forbidden_for_a_reviewer_and_needs_post(web_client, coordinator, reviewer, competition):
    broadcast = _waiting(competition)

    web_client.force_login(reviewer.user)
    assert web_client.post(f"/coordinator/messages/{broadcast.pk}/cancel/").status_code == 403
    web_client.force_login(coordinator)
    assert web_client.get(f"/coordinator/messages/{broadcast.pk}/cancel/").status_code == 405

    broadcast.refresh_from_db()
    assert broadcast.status == BroadcastStatus.SCHEDULED


def test_cancel_after_sending_only_informs(web_client, coordinator, competition):
    broadcast = _waiting(competition)
    MessageBroadcast.objects.filter(pk=broadcast.pk).update(status=BroadcastStatus.SENT)
    web_client.force_login(coordinator)

    response = web_client.post(f"/coordinator/messages/{broadcast.pk}/cancel/", follow=True)

    assert "nie da się już anulować" in response.content.decode()
    broadcast.refresh_from_db()
    assert broadcast.status == BroadcastStatus.SENT


def test_scheduled_list_hides_broadcasts_of_another_competition(
    web_client, coordinator, competition, other_competition
):
    _waiting(competition, subject="Nasz plan")
    _waiting(other_competition, subject="Cudzy plan")
    web_client.force_login(coordinator)

    content = web_client.get("/coordinator/messages/").content.decode()

    assert "Nasz plan" in content
    assert "Cudzy plan" not in content


def test_history_shows_the_send_time(web_client, coordinator, competition):
    from datetime import datetime, timedelta
    from zoneinfo import ZoneInfo

    from django.utils import timezone

    day = (timezone.localtime() - timedelta(days=1)).date()
    broadcast = _waiting(competition, subject="Wysłany z terminem")
    MessageBroadcast.objects.filter(pk=broadcast.pk).update(
        status=BroadcastStatus.SENT,
        scheduled_for=datetime(day.year, day.month, day.day, 6, 17, tzinfo=ZoneInfo("Europe/Warsaw")),
    )
    web_client.force_login(coordinator)

    content = web_client.get("/coordinator/messages/").content.decode()

    assert "Wysłany z terminem" in content
    assert "06:17" in content
