"""Sumy, progi, snapshoty i eksporty z ocenami dziesiętnymi (wydanie 0.35.0).

Cztery rzeczy do udowodnienia:

- **suma jest dokładna** – ``Decimal`` od kolumny do tabeli, bez ``int()`` i bez ``float``: 4,25 + 3,5
  to 7,75, a nie 7,
- **zaokrąglenie sumy ważonej zapada raz, połówka w górę** – do 0,01 w etapie z dowolnymi wartościami
  i (bez zmian) do pełnego punktu w etapie „tylko ze skali”,
- **snapshot jest JSON-em zgodnym wstecz** – liczba całkowita zostaje ``int``-em, ułamkowa staje się
  liczbą JSON, a snapshot zapisany przed wydaniem (same ``int``) renderuje się jak dotąd,
- **CSV niesie kropkę**, XLSX – liczbę.
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal

import pytest

from apps.accounts.tests.factories import CoordinatorFactory, ParticipantFactory
from apps.competitions.models import QualificationMode, StageEntry, StageEntryStatus
from apps.core.exports import csv_response, stage_results_dataset, xlsx_response
from apps.results.models import Anonymization, ResultsPublication
from apps.results.services import (
    apply_qualification,
    compute_stage_results,
    publish_results,
    results_for_participant,
)
from apps.results.statistics import stage_statistics

from .conftest import graded_entry, make_stage, stage_problems

pytestmark = pytest.mark.django_db


def free(stage):
    scale = stage.scoring_scale
    scale.free_values = True
    scale.save(update_fields=["free_values"])
    return stage


def enable_weights(stage):
    competition = stage.edition.competition
    competition.feature_flags = {**(competition.feature_flags or {}), "weighted_scoring": True}
    competition.save(update_fields=["feature_flags"])


def set_weights(problem, numerator, denominator):
    problem.weight_numerator = numerator
    problem.weight_denominator = denominator
    problem.save(update_fields=["weight_numerator", "weight_denominator"])


def test_suma_ocen_ulamkowych_jest_dokladna():
    stage = free(make_stage(problems=2))
    entry = graded_entry(stage, [Decimal("4.25"), Decimal("3.5")])

    rows = compute_stage_results(stage)

    assert rows[0]["points"] == {"1": Decimal("4.25"), "2": Decimal("3.50")}
    assert rows[0]["total"] == Decimal("7.75")
    entry.refresh_from_db()
    assert entry.total_points == Decimal("7.75")


def test_etap_skali_liczy_dokladnie_jak_dotad():
    stage = make_stage(problems=2)
    graded_entry(stage, [6, 2])

    rows = compute_stage_results(stage)

    assert rows[0]["total"] == 8
    assert rows[0]["points"] == {"1": 6, "2": 2}


def test_suma_wazona_w_trybie_dowolnym_zaokragla_do_setnej_polowka_w_gore():
    stage = free(make_stage(problems=1))
    enable_weights(stage)
    set_weights(stage_problems(stage)[0], 1, 2)
    graded_entry(stage, [Decimal("4.25")])

    # 4,25 × 1/2 = 2,125 → 2,13 (ROUND_HALF_UP do 0,01), a nie 2 (pełne punkty) ani 2,12 (bankierskie).
    assert compute_stage_results(stage)[0]["total"] == Decimal("2.13")


def test_trzy_trzecie_w_trybie_dowolnym_zaokraglane_raz_na_koncu():
    stage = free(make_stage(problems=3))
    enable_weights(stage)
    for problem in stage_problems(stage):
        set_weights(problem, 1, 3)
    graded_entry(stage, [Decimal("4.25"), Decimal("4.25"), Decimal("4.26")])

    # (4,25 + 4,25 + 4,26) / 3 = 4,2533… → 4,25; zaokrąglanie po drodze dałoby 1,42+1,42+1,42 = 4,26.
    assert compute_stage_results(stage)[0]["total"] == Decimal("4.25")


def test_suma_wazona_w_trybie_skali_nadal_do_pelnego_punktu():
    stage = make_stage(problems=1)
    enable_weights(stage)
    set_weights(stage_problems(stage)[0], 1, 2)
    graded_entry(stage, [5])

    assert compute_stage_results(stage)[0]["total"] == Decimal(3)


def test_prog_kwalifikacji_z_ulamkiem():
    stage = free(make_stage(problems=2, mode=QualificationMode.MIN_POINTS, min_points=Decimal("7.5")))
    above = graded_entry(stage, [Decimal("4.25"), Decimal("3.5")])
    below = graded_entry(stage, [Decimal("4.25"), Decimal("3")])
    exact = graded_entry(stage, [Decimal("4.5"), Decimal("3")])

    apply_qualification(stage, actor=CoordinatorFactory())

    statuses = dict(StageEntry.objects.filter(stage=stage).values_list("pk", "status"))
    assert statuses[above.pk] == StageEntryStatus.QUALIFIED
    assert statuses[exact.pk] == StageEntryStatus.QUALIFIED
    assert statuses[below.pk] == StageEntryStatus.NOT_QUALIFIED


def test_remis_na_ulamku_to_to_samo_miejsce():
    stage = free(make_stage(problems=1))
    graded_entry(stage, [Decimal("4.25")])
    graded_entry(stage, [Decimal("4.25")])
    graded_entry(stage, [Decimal("4.26")])

    ranks = sorted(row["rank"] for row in compute_stage_results(stage))

    assert ranks == [1, 2, 2]


def test_snapshot_ulamki_jako_liczby_json_a_calkowite_jako_int():
    stage = free(make_stage(problems=2))
    graded_entry(stage, [Decimal("4.25"), Decimal("6")])

    publication = publish_results(stage, CoordinatorFactory(), Anonymization.CODE)

    publication.refresh_from_db()
    row = publication.snapshot[0]
    assert row["points"] == {"1": 4.25, "2": 6}
    assert isinstance(row["points"]["2"], int)
    assert row["total"] == 10.25
    assert list(publication.entry_totals.values()) == [10.25]


def test_snapshot_etapu_skali_bez_zmian_same_int():
    stage = make_stage(problems=2)
    graded_entry(stage, [6, 2])

    publication = publish_results(stage, CoordinatorFactory(), Anonymization.CODE)

    publication.refresh_from_db()
    row = publication.snapshot[0]
    assert row["points"] == {"1": 6, "2": 2}
    assert all(isinstance(value, int) for value in row["points"].values())
    assert isinstance(row["total"], int)


def test_uczestnik_nie_widzi_rozjazdu_gdy_suma_ulamkowa_sie_zgadza(competition):
    stage = free(make_stage(problems=2))
    participant = ParticipantFactory()
    graded_entry(stage, [Decimal("4.25"), Decimal("3.1")], participant=participant)
    publish_results(stage, CoordinatorFactory(), Anonymization.CODE)

    results = results_for_participant(participant.user, competition)

    assert results[0]["total_points"] == Decimal("7.35")
    assert results[0]["published_total"] == Decimal("7.35")
    assert results[0]["differs_from_published"] is False


def test_statystyki_histogram_z_ulamkami_i_srednia_bez_floata():
    stage = free(make_stage(problems=1))
    graded_entry(stage, [Decimal("4.25")])
    graded_entry(stage, [Decimal("4.25")])
    graded_entry(stage, [Decimal("5")])
    publication = publish_results(stage, CoordinatorFactory(), Anonymization.CODE)
    publication = ResultsPublication.objects.select_related("stage", "stage__edition").get(pk=publication.pk)

    stats = stage_statistics(publication)

    bars = stats["problems"][0]["bars"]
    assert [(bar["value"], bar["count"]) for bar in bars] == [(4.25, 2), (5, 1)]
    assert stats["max_total"] == Decimal(5)
    assert stats["mean"] == Decimal("4.5")


def test_csv_z_kropka_a_xlsx_z_liczba():
    stage = free(make_stage(problems=2))
    graded_entry(stage, [Decimal("4.25"), Decimal("3.5")])

    response = csv_response(stage_results_dataset(stage))
    body = b"".join(response.streaming_content).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body), delimiter=";"))
    header, row = rows[0], rows[1]

    assert row[header.index("zadanie 1")] == "4.25"
    assert row[header.index("zadanie 2")] == "3.5"
    assert row[header.index("razem")] == "7.75"

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(xlsx_response(stage_results_dataset(stage)).content))
    sheet = workbook.active
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert float(values[header.index("razem")]) == 7.75


def test_csv_etapu_skali_bez_zer_po_przecinku():
    stage = make_stage(problems=2)
    graded_entry(stage, [6, 2])

    response = csv_response(stage_results_dataset(stage))
    body = b"".join(response.streaming_content).decode("utf-8-sig")
    header, row = list(csv.reader(io.StringIO(body), delimiter=";"))[:2]

    assert row[header.index("razem")] == "8"
    assert row[header.index("zadanie 1")] == "6"


def test_publiczna_tabela_ze_starego_snapshotu_i_z_ulamkami(client_for, competition):
    stage = free(make_stage(problems=2))
    ResultsPublication.objects.create(
        stage=stage,
        anonymization=Anonymization.CODE,
        snapshot=[
            {"rank": 1, "display": "OLM-A", "points": {"1": 6, "2": 5}, "total": 11, "qualified": True},
            {
                "rank": 2,
                "display": "OLM-B",
                "points": {"1": 4.25, "2": 3.5},
                "total": 7.75,
                "qualified": False,
            },
        ],
        entry_totals={},
    )
    stage.results_published_at = stage.opens_at
    stage.save(update_fields=["results_published_at"])

    response = client_for(competition).get(f"/results/{stage.pk}/")

    assert response.status_code == 200
    html = response.content.decode()
    assert '<td class="num total">11</td>' in html
    assert '<td class="num total">7,75</td>' in html
    assert '<td class="num">4,25</td>' in html
    assert "11,00" not in html
    assert "(max 6)" in html
    assert "(max 12)" in html  # suma maksimów dwóch zadań 0–6
