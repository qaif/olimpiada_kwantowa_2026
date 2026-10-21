"""„Postęp klasy” i eksport (UX-PANEL-02 § 5): kto jest na liście, czego nie widać, filtry, koszt."""

from __future__ import annotations

import csv
import io
from datetime import timedelta

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from apps.accounts.profile import anonymise_account
from apps.competitions.models import Problem, StageEntryStatus, StageKind
from apps.competitions.tests.factories import StageFactory
from apps.core.models import AuditLog
from apps.quiz.tests.factories import QuizAttemptFactory, QuizFactory, QuizStageFactory
from apps.results.models import ResultsPublication
from apps.submissions.models import SubmissionStatus
from apps.submissions.tests.factories import SubmissionFactory

from .. import services
from .conftest import TEACHER_EMAIL, enter, flag, make_student, make_teacher

pytestmark = pytest.mark.django_db

PROGRESS = "/supervisor/progress/"


def body(client, path=PROGRESS, **params) -> str:
    response = client.get(path, params)
    assert response.status_code == 200, path
    return response.content.decode()


def publish(stage):
    """Publikacja w mocy: rekord i znacznik na etapie (kształt z ``publish_results``)."""
    now = timezone.now() - timedelta(minutes=1)
    ResultsPublication.objects.update_or_create(
        stage=stage, defaults={"published_at": now, "snapshot": [], "entry_totals": {}}
    )
    stage.results_published_at = now
    stage.save(update_fields=["results_published_at"])


# --- dostęp ---------------------------------------------------------------------------------------


def test_only_students_who_named_this_teacher_and_only_from_this_competition(
    teacher_client, competition, other_competition, stage
):
    mine = make_student(competition, email=TEACHER_EMAIL.upper())
    enter(stage, mine)
    stranger = make_student(competition, email="ktos.inny@szkola.test")
    enter(stage, stranger)
    # Ten sam adres nauczyciela, ale w innym konkursie – profil opiekuna jest per konkurs.
    elsewhere = make_student(other_competition)

    html = body(teacher_client)

    assert mine.public_code in html
    assert stranger.public_code not in html
    assert elsewhere.public_code not in html


def test_anonymised_student_disappears(teacher_client, competition, stage):
    leaves = make_student(competition)
    enter(stage, leaves)
    anonymise_account(leaves.user)

    assert leaves.public_code not in body(teacher_client)


def test_other_roles_get_403_and_anonymous_goes_to_login(client_for, competition, stage):
    participant = make_student(competition)
    client = client_for(competition)
    for url in (PROGRESS, "/supervisor/progress/export.csv", "/supervisor/school/"):
        response = client.get(url)
        assert response.status_code == 302 and "/login/" in response["Location"], url
    client.force_login(participant.user)
    for url in (PROGRESS, "/supervisor/progress/export.csv", "/supervisor/school/"):
        assert client.get(url).status_code == 403, url


def test_teacher_of_another_competition_sees_nothing_here(client_for, competition, other_competition, stage):
    """Profil opiekuna w konkursie B nie jest profilem w A – pod domeną A to obca rola (403)."""
    outsider = make_teacher(other_competition, email="obcy@szkola.test")
    enter(stage, make_student(competition, email="obcy@szkola.test"))
    client = client_for(competition)
    client.force_login(outsider.user)

    assert client.get(PROGRESS).status_code == 403


# --- czego nie widać --------------------------------------------------------------------------------


def test_no_points_and_no_email_before_results(teacher_client, competition, stage):
    student = make_student(competition)
    entry = enter(stage, student)
    entry.total_points = 42
    entry.save(update_fields=["total_points"])
    SubmissionFactory(entry=entry, problem=stage.problems.first(), status=SubmissionStatus.IN_REVIEW)

    html = body(teacher_client)

    assert "w ocenie" in html
    assert "oddane 1 z 2" in html
    assert "Suma punktów" not in html
    assert student.user.email not in html
    progress = services.class_progress(make_teacher_from(student), competition)
    assert progress.rows[0].cells[0].points is None


def make_teacher_from(student):
    from apps.accounts.models import SchoolSupervisor

    return SchoolSupervisor.objects.get(user__email=student.supervisor_email.lower())


def test_points_and_decision_only_after_publication_and_hidden_again_after_withdrawal(
    teacher_client, competition, stage
):
    student = make_student(competition)
    entry = enter(stage, student, status=StageEntryStatus.QUALIFIED)
    entry.total_points = 7
    entry.save(update_fields=["total_points"])
    publish(stage)

    html = body(teacher_client)
    assert "Suma punktów" in html and "zakwalifikowany" in html

    stage.results_published_at = None
    stage.save(update_fields=["results_published_at"])
    html = body(teacher_client)
    assert "Suma punktów" not in html and "zakwalifikowany" not in html


def test_infected_work_is_not_a_submission_and_least_advanced_work_labels_the_stage(
    teacher_client, competition, stage
):
    student = make_student(competition)
    entry = enter(stage, student)
    first, second = stage.problems.order_by("number")
    SubmissionFactory(entry=entry, problem=first, status=SubmissionStatus.FINAL)
    SubmissionFactory(entry=entry, problem=second, status=SubmissionStatus.SUBMITTED, version=1)
    SubmissionFactory(entry=entry, problem=second, status=SubmissionStatus.REJECTED_INFECTED, version=2)

    cell = services.class_progress(make_teacher_from(student), competition).rows[0].cells[0]

    assert cell.submitted == 2
    assert cell.step == "submitted"


def test_quiz_counts_a_finished_attempt(teacher, competition, edition):
    quiz_stage = QuizStageFactory(edition=edition, kind=StageKind.ELIM, name="Test")
    quiz = QuizFactory(stage=quiz_stage)
    student = make_student(competition)
    entry = enter(quiz_stage, student)
    QuizAttemptFactory(quiz=quiz, entry=entry, submitted_at=timezone.now())

    row = services.class_progress(teacher, competition).rows[0]

    assert row.cells[0].submitted == 1 and row.cells[0].total == 1
    assert row.missing_work is False


def test_training_stage_is_not_a_column(teacher, competition, edition, stage):
    StageFactory(edition=edition, kind=StageKind.TRAINING, name="Trening")
    enter(stage, make_student(competition))

    progress = services.class_progress(teacher, competition)

    assert [column.stage.name for column in progress.columns] == ["Etap I"]


# --- flaga formalności ---------------------------------------------------------------------------


def test_details_are_not_read_or_shown_without_the_flag(teacher_client, teacher, competition, stage):
    flag(competition, "student_status_certificate")
    enter(stage, make_student(competition))

    html = body(teacher_client)
    row = services.class_progress(teacher, competition).rows[0]

    assert "Zgoda rodzica" not in html and "Zaświadczenie" not in html
    assert row.guardian is None and row.certificate is None and row.workshops is None


def test_details_with_the_flag(teacher_client, teacher, competition, edition, stage):
    from apps.student_status.models import CertificateStatus, StudentStatusCertificate

    flag(competition, "student_status_certificate")
    flag(competition, "supervisor_class_details")
    minor = make_student(competition, guardian_consent=False, user__last_name="Adamska")
    adult = make_student(competition, birth_year=2000, user__last_name="Borowski")
    pending = make_student(competition, guardian_consent=True, user__last_name="Cichy")
    for student in (minor, adult, pending):
        enter(stage, student)
    StudentStatusCertificate.objects.create(
        participant=pending,
        edition=edition,
        version=1,
        status=CertificateStatus.PENDING,
        mime="application/pdf",
        size_bytes=10,
    )

    rows = {row.participant.pk: row for row in services.class_progress(teacher, competition).rows}
    html = body(teacher_client)

    assert rows[minor.pk].guardian == services.GUARDIAN_MISSING
    assert rows[adult.pk].guardian == services.GUARDIAN_NA
    assert rows[pending.pk].guardian == services.GUARDIAN_OK
    assert rows[minor.pk].certificate == services.CERT_MISSING
    assert rows[pending.pk].certificate == services.CERT_PENDING
    assert "Zgoda rodzica" in html and "czeka na weryfikację" in html
    only_missing = body(teacher_client, zgoda="missing")
    assert minor.public_code in only_missing and adult.public_code not in only_missing


# --- filtry i sortowanie ---------------------------------------------------------------------------


def test_filters_grade_account_and_missing_work(teacher_client, competition, stage):
    worker = make_student(competition, grade=2, user__last_name="Pracowity")
    SubmissionFactory(entry=enter(stage, worker), problem=stage.problems.first())
    idle = make_student(competition, grade=3, user__last_name="Leniwy")
    enter(stage, idle)
    invited = make_student(competition, grade=3, user__last_name="Zaproszony", invited_at=timezone.now())
    invited.user.email_verified_at = None
    invited.user.save(update_fields=["email_verified_at"])
    disqualified = make_student(competition, grade=3, user__last_name="Wykluczony")
    enter(stage, disqualified, status=StageEntryStatus.DISQUALIFIED)

    by_grade = body(teacher_client, klasa="2")
    assert worker.public_code in by_grade and idle.public_code not in by_grade

    missing = body(teacher_client, praca="missing")
    assert idle.public_code in missing
    assert worker.public_code not in missing
    assert disqualified.public_code not in missing  # zdyskwalifikowany nie „zalega”

    by_account = body(teacher_client, konto="invited")
    assert invited.public_code in by_account and worker.public_code not in by_account

    nonsense = body(teacher_client, klasa="99", konto="x", sort="hack")
    assert worker.public_code in nonsense and idle.public_code in nonsense


def test_sorting(teacher, competition, stage):
    a = make_student(competition, user__last_name="Adamczyk", grade=3)
    b = make_student(competition, user__last_name="Bąk", grade=1)
    SubmissionFactory(entry=enter(stage, b), problem=stage.problems.first())
    progress = services.class_progress(teacher, competition)

    def order(**params):
        filters = services.parse_filters(params, progress)
        return [row.participant.pk for row in services.apply_filters(progress, filters)]

    assert order() == [a.pk, b.pk]
    assert order(sort="-name") == [b.pk, a.pk]
    assert order(sort="grade") == [b.pk, a.pk]
    assert order(sort="-work") == [b.pk, a.pk]


# --- eksport ---------------------------------------------------------------------------------------


def test_csv_has_the_screen_columns_no_email_and_is_audited(teacher_client, teacher, competition, stage):
    student = make_student(competition, user__first_name="Ola", user__last_name="Nowak")
    SubmissionFactory(entry=enter(stage, student), problem=stage.problems.first())
    make_student(competition, user__last_name="Filtrowany", grade=1)

    response = teacher_client.get("/supervisor/progress/export.csv", {"klasa": "3"})

    assert response.status_code == 200
    text = b"".join(response.streaming_content).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    assert rows[0][:4] == ["Imię i nazwisko", "Kod uczestnika", "Klasa", "Konto"]
    assert "Etap I: oddane prace" in rows[0]
    assert "Etap I: punkty" in rows[0]
    assert len(rows) == 2  # filtr klasy z ekranu działa też w pliku
    assert rows[1][:2] == ["Ola Nowak", student.public_code]
    assert "1/2" in rows[1]
    assert student.user.email not in text
    entry = AuditLog.objects.get(action="export.generated", target_id=str(teacher.pk))
    assert entry.diff["kind"] == "supervisor_class_progress"
    assert entry.diff["format"] == "csv" and entry.diff["rows"] == 1


def test_xlsx_and_unknown_format(teacher_client, competition, stage):
    from openpyxl import load_workbook

    enter(stage, make_student(competition))

    response = teacher_client.get("/supervisor/progress/export.xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [cell.value for cell in sheet[1]][:2] == ["Imię i nazwisko", "Kod uczestnika"]
    assert sheet.max_row == 2
    assert teacher_client.get("/supervisor/progress/export.pdf").status_code == 404


def test_export_links_carry_the_filters(teacher_client, competition, stage):
    enter(stage, make_student(competition))

    html = body(teacher_client, praca="missing", sort="-name")

    assert reverse("web:supervisor-progress-export", args=["csv"]) + "?praca=missing&amp;sort=-name" in html


# --- koszt ---------------------------------------------------------------------------------------


def _queries(client, path) -> int:
    with CaptureQueriesContext(connection) as captured:
        response = client.get(path)
        assert response.status_code == 200, path
        if response.streaming:
            b"".join(response.streaming_content)
    return len(captured)


COST_PATHS = (PROGRESS, "/supervisor/", "/supervisor/progress/export.csv")


def test_query_count_does_not_grow_with_the_class(teacher_client, competition, edition, stage):
    """Stała liczba zapytań: 2 i 6 uczniów (prace, test online, flaga formalności) – to samo (D9).

    Ekran, „Start” i eksport – wszystkie trzy chodzą przez ``class_progress``.
    """
    flag(competition, "supervisor_class_details")
    flag(competition, "student_status_certificate")
    quiz_stage = QuizStageFactory(edition=edition, kind=StageKind.DISTRICT, name="Test")
    quiz = QuizFactory(stage=quiz_stage)

    def add(count):
        for _index in range(count):
            student = make_student(competition)
            entry = enter(stage, student)
            SubmissionFactory(entry=entry, problem=Problem.objects.filter(stage=stage).first())
            QuizAttemptFactory(quiz=quiz, entry=enter(quiz_stage, student), submitted_at=timezone.now())

    add(2)
    for path in COST_PATHS:
        _queries(teacher_client, path)  # rozgrzewka pamięci procesu (przełączniki, strona „Warsztaty”)
    small = {path: _queries(teacher_client, path) for path in COST_PATHS}
    add(4)
    large = {path: _queries(teacher_client, path) for path in COST_PATHS}

    assert large == small
