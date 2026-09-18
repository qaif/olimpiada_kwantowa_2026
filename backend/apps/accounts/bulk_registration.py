"""Hurtowe zapraszanie uczniów: nauczyciel wgrywa listę klasy, konto zakłada uczeń.

Skąd ta funkcja. Zgłoszenie z konsultacji ze szkołami brzmiało prosto: „mam dwudziestu uczniów,
nie będę dwudziestu razy tłumaczyć, jak się zarejestrować”. Naturalna odpowiedź – „nauczyciel
zakłada konta i rozdaje hasła” – jest jednak wykluczona, i to nie ze względów wygody:

- **zgody są oświadczeniem ucznia** (regulamin, RODO, a dla niepełnoletnich zgoda opiekuna).
  Nauczyciel nie może ich złożyć w cudzym imieniu, więc konto zrobione „pod klucz” byłoby kontem
  bez ani jednej ważnej zgody,
- **hasło rozdane przez osobę trzecią nie jest poświadczeniem.** Uczeń, którego hasło zna
  nauczyciel, nie ma konta – ma konto współdzielone.

Dlatego import dzieli się na dwie czynności wykonane przez dwie różne osoby:

1. **nauczyciel wgrywa listę** (imię, nazwisko, e-mail, rocznik, klasa; telefon i adres opiekuna
   prawnego opcjonalnie). Powstaje konto w stanie „zaproszony”: nieaktywne, z hasłem
   **nieużywalnym** (``set_unusable_password``), z profilem uczestnika wskazującym szkołę
   nauczyciela i jego adres jako adres opiekuna szkolnego. Takie konto nie zaloguje się żadną
   drogą – ani hasłem (nie ma go), ani przez dostawcę OAuth (adres zajęty przez konto bez
   powiązania kończy się odmową w ``apps.accounts.adapters``),
2. **uczeń przyjmuje zaproszenie** pod podpisanym linkiem (``/zaproszenie/<token>/``, 14 dni):
   ustawia własne hasło, uzupełnia dane, których nauczyciel nie miał prawa znać za niego,
   i **sam** składa komplet zgód. Dopiero to aktywuje konto – tą samą funkcją
   (``apps.accounts.activation.mark_activated``), co kliknięcie linku aktywacyjnego.

Profil uczestnika powstaje więc **przed** zgodami i jest to jedyne miejsce w systemie, gdzie tak
się dzieje. Pole ``Participant.gdpr_consent_at`` jest z tego powodu nullowalne, a puste znaczy
dokładnie „ten człowiek jeszcze niczego nie oświadczył”. Reguła „bez zgody nie ma udziału
w zawodach” nie słabnie: konto bez zgód nie przechodzi aktywacji, więc się nie zaloguje, nie
zapisze do etapu i niczego nie odda. Zyskujemy natomiast to, po co cały import istnieje –
nauczyciel od pierwszej chwili widzi listę swoich uczniów i to, kto zaproszenie przyjął.

Czego import **nie** robi:

- **nie zakłada kont, które już istnieją.** Adres zajęty przez konto uczestnika jest
  *dowiązywany*: w profilu tego uczestnika staje się adres opiekuna szkolnego i tyle. Nauczyciel
  dostaje o tym notkę w podglądzie, bo inaczej wyglądałoby to na cichy błąd,
- **nie zbiera zgód i nie ustawia haseł.** Patrz wyżej,
- **nie wpisuje danych osobowych do audytu.** Wpisy niosą liczby i identyfikatory wierszy; kogo
  dotyczą, mówi ``target_id``. Audyt czytają osoby, które nie mają wglądu w listy klasowe.

Podgląd przed zapisem jest częścią funkcji, a nie ozdobą: plik z arkusza szkolnego zawiera
literówki w adresach, puste wiersze, stopki i uczniów już zarejestrowanych. Zapis bez pokazania,
co dokładnie się wydarzy, znaczyłby dwadzieścia listów wysłanych pod przypadkowe adresy.
"""

from __future__ import annotations

import csv
import io
import logging
import unicodedata
from dataclasses import dataclass, field

from django.core import signing
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import EmailValidator
from django.db import transaction
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.utils.translation import gettext_lazy
from rest_framework import status

from apps.core.api import DomainError
from apps.core.models import audit

from .activation import absolute_url, mark_activated, queue_mail
from .consents import is_minor
from .models import GROUP_PARTICIPANT, MAX_GRADE, MIN_GRADE, Participant, User
from .phones import normalize_phone
from .supervisors import normalize_supervisor_email

logger = logging.getLogger(__name__)

#: Ile wierszy przyjmujemy z jednego pliku. Limit jest decyzją organizatora, a nie granicą
#: techniczną: import obsługuje **klasę albo szkołę**, a plik z tysiącami adresów znaczy, że ktoś
#: wkleił cudzą bazę. Pięćset mieści największe zespoły szkół z zapasem.
MAX_ROWS = 500

#: Górna granica rozmiaru pliku. Arkusz z pięciuset wierszami waży kilkadziesiąt kilobajtów;
#: dwa megabajty zostawiają zapas na formatowanie XLSX i zatrzymują wgranie czegoś zupełnie innego,
#: zanim openpyxl zacznie to rozpakowywać.
MAX_UPLOAD_BYTES = 2 * 1024 * 1024

#: Sól podpisu linku zaproszenia. Osobna od aktywacji konta, zmiany adresu i zgody opiekuna –
#: to czwarte, inne uprawnienie („ustaw hasło i złóż zgody za to konto”) i token jednego z nich
#: nie może zadziałać w miejscu drugiego.
INVITE_SALT = "apps.accounts.student-invite"

#: Ważność linku zaproszenia – czternaście dni, tyle samo, co prośba o zgodę opiekuna. Cztery
#: godziny (jak przy aktywacji) byłyby tu nieporozumieniem: tam okno pilnuje **własnej**
#: rejestracji sprzed chwili, a tu list przychodzi znienacka, bywa w szkolnej skrzynce czytanej
#: raz w tygodniu, a konto na nikogo nie czeka i niczego nie blokuje.
INVITE_MAX_AGE = 14 * 24 * 3600
INVITE_DAYS = INVITE_MAX_AGE // (24 * 3600)

#: Sól podpisu koszyka wierszy niesionego między podglądem a zatwierdzeniem (patrz ``pack_rows``).
PREVIEW_SALT = "apps.accounts.student-import-preview"

#: Ile czasu ma nauczyciel na zatwierdzenie podglądu. Dwie godziny: tyle, żeby dało się odejść od
#: komputera na lekcję, i na tyle mało, żeby zatwierdzić nie dało się listy sprzed tygodnia.
PREVIEW_MAX_AGE = 2 * 3600

INVITE_SUBJECT_TEMPLATE = "registration/invite_student_subject.txt"
INVITE_BODY_TEMPLATE = "registration/invite_student_body.txt"

INVALID_TOKEN_MESSAGE = gettext_lazy("Link z zaproszeniem jest nieprawidłowy albo wygasł.")


# --- kolumny pliku --------------------------------------------------------------------------


@dataclass(frozen=True)
class Column:
    """Jedna kolumna listy: klucz w kodzie, nagłówek dla człowieka i jego dopuszczalne warianty.

    Warianty są potrzebne, bo plik przychodzi z arkusza prowadzonego przez nauczyciela, a nie
    z naszego eksportu: „Imię”, „imie”, „IMIONA” i „Rok urodzenia” obok „rocznik” to ten sam
    zamiar. Dopasowanie idzie po postaci **bez znaków diakrytycznych i bez interpunkcji**
    (``_header_key``), więc „E-mail”, „e mail” i „EMAIL” trafiają w ten sam klucz.
    """

    key: str
    label: str
    aliases: tuple[str, ...]
    required: bool = True


def _header_key(text: str) -> str:
    """Postać porównawcza nagłówka: małe litery ASCII, bez interpunkcji i spacji."""
    folded = unicodedata.normalize("NFKD", (text or "").strip().lower())
    stripped = "".join(char for char in folded if not unicodedata.combining(char))
    # Polskie „ł” nie rozkłada się na literę i znak łączący – po NFKD zostaje samo „ł”.
    stripped = stripped.replace("ł", "l")
    return "".join(char for char in stripped if char.isalnum())


COLUMNS: tuple[Column, ...] = (
    Column("first_name", "imię", ("imie", "imiona", "imieucznia")),
    Column("last_name", "nazwisko", ("nazwisko", "nazwiskoucznia")),
    Column("email", "e-mail", ("email", "adresemail", "emailucznia", "adresemailucznia")),
    Column("birth_year", "rok urodzenia", ("rokurodzenia", "rocznik")),
    Column("grade", "klasa", ("klasa",)),
    Column("phone", "telefon", ("telefon", "nrtelefonu", "numertelefonu"), required=False),
    Column(
        "guardian_email",
        "e-mail opiekuna prawnego",
        ("emailopiekunaprawnego", "adresemailopiekunaprawnego", "emailrodzica", "opiekunprawny"),
        required=False,
    ),
)

#: Kolumna dostępna wyłącznie w imporcie koordynatora: pozwala przypiąć ucznia do **cudzego**
#: opiekuna szkolnego. W imporcie nauczyciela byłaby bez sensu (adres opiekuna jest z definicji
#: jego własny) i groźna: dawałaby dowolnemu nauczycielowi przypisanie uczniów komu innemu.
SUPERVISOR_COLUMN = Column(
    "supervisor_email",
    "e-mail opiekuna szkolnego",
    ("emailopiekunaszkolnego", "adresemailopiekunaszkolnego", "opiekunszkolny", "emailnauczyciela"),
    required=False,
)


#: Kolumny, które dokłada **konfiguracja konkursu**, a nie droga importu (§ 1.3.4, § 1.2.4, § 1.4).
#: Każda jest nieobowiązkowa i każda pojawia się wyłącznie razem ze swoją flagą – plik
#: przygotowany wcześniej wczytuje się dalej bez zmiany, a instrukcja na ekranie nie wymienia
#: kolumn, których ten konkurs i tak nie przyjmie.
REGION_COLUMN = Column("region", "region", ("region", "wojewodztwo", "okreg"), required=False)
CATEGORY_COLUMN = Column("category", "kategoria", ("kategoria", "kodkategorii"), required=False)
INSTITUTION_TYPE_COLUMN = Column(
    "institution_type",
    "typ placówki",
    ("typplacowki", "rodzajplacowki", "typinstytucji"),
    required=False,
)
INSTITUTION_ID_COLUMN = Column(
    "custom_institution_id",
    "placówka (identyfikator)",
    ("placowkaidentyfikator", "placowkaid", "idplacowki", "identyfikatorplacowki"),
    required=False,
)
INSTITUTION_NAME_COLUMN = Column(
    "institution_name", "placówka", ("placowka", "nazwaplacowki"), required=False
)
COUNTRY_COLUMN = Column("country", "kraj", ("kraj", "panstwo", "kodkraju"), required=False)

#: Flaga kategorii z katalogu § 0.6. Nazwa jest tu przepisana, a nie zaimportowana z
#: ``apps.results.services``: warstwa kont nie zna warstwy wyników i poznać jej nie ma – import
#: w tę stronę zamknąłby pętlę między aplikacjami. Jedynym odczytem zostaje ``has_feature``.
CATEGORIES_FLAG = "categories"


def extra_columns(competition=None) -> tuple[Column, ...]:
    """Kolumny dołożone przez konkurs – pusta krotka, dopóki żadna flaga nie jest włączona.

    Trzy powody, dla których to jest jedno miejsce, a nie warunek w każdym z trzech ekranów:

    - **instrukcja i parser mają widzieć ten sam zestaw.** Kolumna wymieniona w opisie pliku,
      której parser nie zna (albo odwrotnie), jest najgorszym rodzajem pomyłki w imporcie:
      nauczyciel wypełnia rubrykę, która nigdzie nie trafia;
    - **Konkurs #1 ma tu dostać dokładnie dzisiejszy zestaw** (§ 0.1). Wszystkie flagi są u niego
      domyślne, więc funkcja oddaje ``()`` i nie dotyka bazy ani razu;
    - ``competition=None`` znaczy „nie wiadomo, w jakim konkursie” i też oddaje ``()``. Odwrót
      jest **miękki** i celowo po stronie dzisiejszego zachowania: wołający spoza żądania (komenda,
      test, zadanie Celery) dostaje kolumny, które umiał obsłużyć przed etapem 2.
    """
    if competition is None:
        return ()
    from .services import CUSTOM_REGIONS_FLAG, REGISTRATION_PROFILE_FLAG, custom_directory_enabled

    extra: list[Column] = []
    if competition.has_feature(CUSTOM_REGIONS_FLAG):
        extra.append(REGION_COLUMN)
    if competition.has_feature(CATEGORIES_FLAG):
        extra.append(CATEGORY_COLUMN)
    if competition.has_feature(REGISTRATION_PROFILE_FLAG):
        extra.extend((INSTITUTION_TYPE_COLUMN, INSTITUTION_NAME_COLUMN, COUNTRY_COLUMN))
    if custom_directory_enabled(competition):
        extra.append(INSTITUTION_ID_COLUMN)
    return tuple(extra)


def columns_for(*, with_supervisor: bool, competition=None) -> tuple[Column, ...]:
    """Zestaw kolumn dla danej drogi importu – nauczyciela albo koordynatora.

    Kolumny konkursu stoją **na końcu**, za kolumną opiekuna szkolnego: kolejność w nagłówku i tak
    nie ma znaczenia (dopasowanie idzie po nazwie), a dzisiejszy początek wiersza zostaje wtedy
    znak w znak taki, jak w arkuszach, które nauczyciele mają już przygotowane.
    """
    base = (*COLUMNS, SUPERVISOR_COLUMN) if with_supervisor else COLUMNS
    return (*base, *extra_columns(competition))


def header_line(*, with_supervisor: bool, competition=None) -> str:
    """Wzorcowy wiersz nagłówka do pokazania na stronie importu (i do skopiowania do arkusza)."""
    return ";".join(
        column.label for column in columns_for(with_supervisor=with_supervisor, competition=competition)
    )


# --- wiersz i jego rozstrzygnięcie ------------------------------------------------------------

#: Co import zrobi z wierszem. Trzy wartości, bo tyle jest różnych zdarzeń – i każde znaczy dla
#: nauczyciela co innego: „założę konto”, „dopiszę się do istniejącego”, „pomijam i oto dlaczego”.
ACTION_CREATE = "create"
ACTION_LINK = "link"
ACTION_SKIP = "skip"


@dataclass
class ImportRow:
    """Jeden wiersz listy po oczyszczeniu, razem z rozstrzygnięciem i powodami.

    ``errors`` blokują wiersz (zostaje pominięty), ``notes`` wyłącznie opisują, co się stanie –
    rozdział jest istotny, bo „uczeń niepełnoletni bez adresu opiekuna prawnego” **nie** jest
    powodem do odrzucenia: zgodę opiekuna zbieramy od ucznia po przyjęciu zaproszenia
    (``apps.accounts.guardian``), a adres w pliku jest wyłącznie udogodnieniem.
    """

    number: int
    first_name: str = ""
    last_name: str = ""
    email: str = ""
    birth_year: int | None = None
    grade: int | None = None
    phone: str = ""
    guardian_email: str = ""
    supervisor_email: str = ""
    #: Surowa treść kolumn dokładanych przez konkurs (:func:`extra_columns`). Puste napisy znaczą
    #: „tej kolumny w pliku nie było” – i to jest stan Konkursu #1 we wszystkich wierszach.
    region_code: str = ""
    category_code: str = ""
    institution_type: str = ""
    custom_institution_id: str = ""
    institution_name: str = ""
    country: str = ""
    action: str = ACTION_CREATE
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    #: Rozstrzygnięcia liczone przez :func:`validate_rows` – **nie** jadą w koszyku i nie wchodzą
    #: do ``payload``: między podglądem a zatwierdzeniem region mógł zostać wygaszony, a kategoria
    #: przestawiona, więc zapis liczy je od nowa (tak samo jak decyzję „założyć czy dowiązać”).
    district: str = ""
    region: object | None = None
    category: object | None = None
    institution: dict | None = None

    @property
    def full_name(self) -> str:
        return f"{self.first_name} {self.last_name}".strip()

    def payload(self) -> dict:
        """Postać niesiona między podglądem a zatwierdzeniem – same dane, bez rozstrzygnięcia.

        Rozstrzygnięcia świadomie **nie** przenosimy: między jednym ekranem a drugim ktoś mógł
        założyć konto na ten adres. Zatwierdzenie liczy je od nowa (``validate_rows``), więc
        koszyk jest listą danych, a nie listą decyzji do wykonania w ciemno.
        """
        data = {
            "n": self.number,
            "f": self.first_name,
            "l": self.last_name,
            "e": self.email,
            "b": self.birth_year,
            "g": self.grade,
            "p": self.phone,
            "gu": self.guardian_email,
            "s": self.supervisor_email,
        }
        # Kolumny konkursu wchodzą do koszyka **wyłącznie wypełnione**. Nie jest to oszczędność
        # bajtów: dzięki temu koszyk pliku bez tych kolumn jest znak w znak taki sam, jak przed
        # etapem 2, więc podgląd Konkursu #1 nie zmienia się nawet w polu ukrytym (§ 0.1).
        data.update(
            {
                key: value
                for key, value in (
                    ("r", self.region_code),
                    ("c", self.category_code),
                    ("it", self.institution_type),
                    ("ii", self.custom_institution_id),
                    ("in", self.institution_name),
                    ("k", self.country),
                )
                if value
            }
        )
        return data

    @classmethod
    def from_payload(cls, data: dict) -> ImportRow:
        return cls(
            number=int(data.get("n") or 0),
            first_name=str(data.get("f") or ""),
            last_name=str(data.get("l") or ""),
            email=str(data.get("e") or ""),
            birth_year=data.get("b"),
            grade=data.get("g"),
            phone=str(data.get("p") or ""),
            guardian_email=str(data.get("gu") or ""),
            supervisor_email=str(data.get("s") or ""),
            region_code=str(data.get("r") or ""),
            category_code=str(data.get("c") or ""),
            institution_type=str(data.get("it") or ""),
            custom_institution_id=str(data.get("ii") or ""),
            institution_name=str(data.get("in") or ""),
            country=str(data.get("k") or ""),
        )


@dataclass
class ImportPreview:
    """Wynik odczytania pliku: wiersze, koszyk do zatwierdzenia i licznik rozstrzygnięć."""

    rows: list[ImportRow]
    token: str

    @property
    def to_create(self) -> int:
        return sum(1 for row in self.rows if row.action == ACTION_CREATE)

    @property
    def to_link(self) -> int:
        return sum(1 for row in self.rows if row.action == ACTION_LINK)

    @property
    def skipped(self) -> int:
        return sum(1 for row in self.rows if row.action == ACTION_SKIP)

    @property
    def has_importable(self) -> bool:
        return bool(self.to_create or self.to_link)


# --- czytanie pliku -------------------------------------------------------------------------


def _decode_csv(data: bytes) -> str:
    """Tekst pliku CSV. Dwa kodowania, bo tyle wychodzi z Excela na komputerze nauczyciela.

    ``utf-8-sig`` zdejmuje BOM, który Excel dokleja przy „CSV UTF-8”. Gdy to nie jest UTF-8,
    zostaje ``cp1250`` – domyślne kodowanie polskiego Excela przy zwykłym „CSV (rozdzielany
    przecinkami)”. Odmowa z komunikatem „zapisz jako UTF-8” byłaby przerzuceniem na nauczyciela
    problemu, którego nie wywołał i którego nazwy nie ma obowiązku znać.
    """
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DomainError(
        "Nie udało się odczytać pliku – zapisz go jako CSV w kodowaniu UTF-8 albo jako XLSX.",
        "IMPORT_ENCODING",
        status.HTTP_400_BAD_REQUEST,
    )


def _csv_table(data: bytes) -> list[list[str]]:
    """Wiersze pliku CSV. Separator rozpoznajemy sami – polski Excel zapisuje średnikami."""
    text = _decode_csv(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        # Plik z jedną kolumną albo pusty: sniffer nie ma czego rozpoznać, a średnik jest
        # wariantem, który wychodzi z polskiego Excela najczęściej.
        delimiter = ";"
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[(cell or "").strip() for cell in row] for row in reader]


def _xlsx_table(data: bytes) -> list[list[str]]:
    """Wiersze pierwszego arkusza XLSX. ``data_only`` – formuły czytamy jako ich wynik."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl podnosi kilkanaście różnych klas
        raise DomainError(
            "Nie udało się otworzyć pliku XLSX. Zapisz go ponownie albo wgraj CSV.",
            "IMPORT_XLSX_BROKEN",
            status.HTTP_400_BAD_REQUEST,
        ) from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        table = []
        for row in sheet.iter_rows(values_only=True):
            table.append(["" if cell is None else str(cell).strip() for cell in row])
        return table
    finally:
        workbook.close()


def read_table(upload) -> list[list[str]]:
    """Surowa tabela z wgranego pliku – CSV albo XLSX, rozpoznawane po nazwie.

    Rozmiar sprawdzamy **przed** czytaniem: openpyxl rozpakowuje archiwum ZIP, więc plik wgrany
    w złej wierze („zip bomb”) urósłby w pamięci procesu, zanim ktokolwiek zobaczyłby nagłówek.
    """
    name = (getattr(upload, "name", "") or "").lower()
    size = getattr(upload, "size", None)
    if size is not None and size > MAX_UPLOAD_BYTES:
        raise DomainError(
            f"Plik jest za duży (maks. {MAX_UPLOAD_BYTES // (1024 * 1024)} MB).",
            "IMPORT_TOO_LARGE",
            status.HTTP_400_BAD_REQUEST,
        )
    data = upload.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise DomainError(
            f"Plik jest za duży (maks. {MAX_UPLOAD_BYTES // (1024 * 1024)} MB).",
            "IMPORT_TOO_LARGE",
            status.HTTP_400_BAD_REQUEST,
        )
    if not data.strip():
        raise DomainError("Plik jest pusty.", "IMPORT_EMPTY", status.HTTP_400_BAD_REQUEST)
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _xlsx_table(data)
    if name.endswith(".csv") or name.endswith(".txt"):
        return _csv_table(data)
    raise DomainError(
        "Obsługujemy pliki CSV i XLSX.",
        "IMPORT_FORMAT",
        status.HTTP_400_BAD_REQUEST,
    )


def map_header(header: list[str], columns: tuple[Column, ...]) -> dict[str, int]:
    """Nagłówek → ``{klucz kolumny: numer kolumny}``. Brak kolumny wymaganej to odmowa.

    Odmowa jest tu lepsza od domyślania się po kolejności: plik bez nagłówka albo z kolumnami
    w innej kolejności dałby listę uczniów, w której nazwiska trafiły do rubryki „e-mail”,
    i dowiedzielibyśmy się o tym z dwudziestu odbitych listów.
    """
    seen = {_header_key(cell): index for index, cell in enumerate(header) if (cell or "").strip()}
    mapping: dict[str, int] = {}
    missing: list[str] = []
    for column in columns:
        for alias in column.aliases:
            if alias in seen:
                mapping[column.key] = seen[alias]
                break
        else:
            if column.required:
                missing.append(column.label)
    if missing:
        raise DomainError(
            "W pliku brakuje kolumn: " + ", ".join(missing) + ".",
            "IMPORT_HEADER",
            status.HTTP_400_BAD_REQUEST,
        )
    return mapping


def _cell(row: list[str], mapping: dict[str, int], key: str) -> str:
    index = mapping.get(key)
    if index is None or index >= len(row):
        return ""
    return (row[index] or "").strip()


def _clean_year(text: str, row: ImportRow) -> int | None:
    """Rocznik z komórki. Arkusz potrafi podać „2008.0” – to nadal jest rok, a nie błąd."""
    value = text.replace(" ", "")
    if value.endswith(".0"):
        value = value[:-2]
    if not value:
        row.errors.append("brak roku urodzenia")
        return None
    try:
        year = int(value)
    except ValueError:
        row.errors.append("rok urodzenia nie jest liczbą")
        return None
    if not 1900 <= year <= timezone.localdate().year:
        row.errors.append("rok urodzenia poza zakresem")
        return None
    return year


def _clean_grade(text: str, row: ImportRow) -> int | None:
    """Klasa 1–5. „1A”, „3 b” i „2.0” znaczą klasę – litera oddziału nas nie interesuje."""
    digits = "".join(char for char in text if char.isdigit())
    if not digits:
        row.errors.append("brak klasy")
        return None
    grade = int(digits[0])
    if not MIN_GRADE <= grade <= MAX_GRADE:
        row.errors.append(f"klasa spoza zakresu {MIN_GRADE}–{MAX_GRADE}")
        return None
    return grade


def _clean_email(text: str, row: ImportRow) -> str:
    """Adres ucznia. Bez niego nie ma do kogo wysłać zaproszenia, więc jest to błąd wiersza."""
    email = (text or "").strip().lower()
    if not email:
        row.errors.append("brak adresu e-mail")
        return ""
    try:
        EmailValidator()(email)
    except DjangoValidationError:
        row.errors.append("adres e-mail jest nieprawidłowy")
        return ""
    return email


def parse_table(
    table: list[list[str]], *, columns: tuple[Column, ...], default_supervisor_email: str = ""
) -> list[ImportRow]:
    """Tabela → wiersze po oczyszczeniu. Bez zapytań do bazy: to jest sama treść pliku.

    Numer wiersza liczymy tak, jak widzi go człowiek w arkuszu – razem z nagłówkiem. Inaczej
    „wiersz 7” z komunikatu wskazywałby w Excelu na kogoś innego.
    """
    if not table:
        raise DomainError("Plik jest pusty.", "IMPORT_EMPTY", status.HTTP_400_BAD_REQUEST)
    mapping = map_header(table[0], columns)
    rows: list[ImportRow] = []
    for offset, raw in enumerate(table[1:], start=2):
        if not any((cell or "").strip() for cell in raw):
            # Puste wiersze na końcu arkusza są normą, a nie błędem – arkusz „ma” tysiąc wierszy.
            continue
        if len(rows) >= MAX_ROWS:
            raise DomainError(
                f"Plik ma więcej niż {MAX_ROWS} wierszy. Podziel listę na mniejsze pliki.",
                "IMPORT_TOO_MANY_ROWS",
                status.HTTP_400_BAD_REQUEST,
            )
        row = ImportRow(number=offset)
        row.first_name = _cell(raw, mapping, "first_name")[:150]
        row.last_name = _cell(raw, mapping, "last_name")[:150]
        if not row.first_name or not row.last_name:
            row.errors.append("brak imienia albo nazwiska")
        row.email = _clean_email(_cell(raw, mapping, "email"), row)
        row.birth_year = _clean_year(_cell(raw, mapping, "birth_year"), row)
        row.grade = _clean_grade(_cell(raw, mapping, "grade"), row)
        phone = _cell(raw, mapping, "phone")
        if phone:
            try:
                row.phone = normalize_phone(phone, required=False)
            except DomainError:
                # Telefon jest udogodnieniem, a nie warunkiem konta: zły numer zostawiamy pusty
                # i mówimy o tym, zamiast odrzucać ucznia z powodu spacji w numerze.
                row.notes.append("numer telefonu pominięty – niepoprawny zapis")
        guardian = _cell(raw, mapping, "guardian_email").lower()
        if guardian:
            guardian_row = ImportRow(number=offset)
            row.guardian_email = _clean_email(guardian, guardian_row)
            if guardian_row.errors:
                row.guardian_email = ""
                row.notes.append("adres opiekuna prawnego pominięty – niepoprawny zapis")
        supervisor = _cell(raw, mapping, "supervisor_email")
        row.supervisor_email = normalize_supervisor_email(supervisor) or default_supervisor_email
        # Kolumny konkursu czytamy **surowo**: rozstrzyga je ``validate_rows``, bo region, kategoria
        # i placówka są wierszami w bazie, a ta funkcja z założenia nie zadaje ani jednego pytania.
        # Gdy kolumny w pliku nie ma, ``_cell`` oddaje pusty napis – czyli dokładnie to, co wiersz
        # ma dziś (§ 0.1).
        row.region_code = _cell(raw, mapping, "region")[:40]
        row.category_code = _cell(raw, mapping, "category")[:32]
        row.institution_type = _cell(raw, mapping, "institution_type").upper()[:32]
        row.custom_institution_id = _cell(raw, mapping, "custom_institution_id")[:32]
        row.institution_name = _cell(raw, mapping, "institution_name")[:255]
        row.country = _cell(raw, mapping, "country").upper()[:32]
        rows.append(row)
    if not rows:
        raise DomainError(
            "Plik nie zawiera ani jednego wiersza z danymi.",
            "IMPORT_NO_ROWS",
            status.HTTP_400_BAD_REQUEST,
        )
    return rows


# --- kolumny konkursu: region, kategoria, placówka ---------------------------------------------


def _column_error(column: Column, value: str, detail) -> str:
    """Komunikat wiersza nazywający **kolumnę i wartość**.

    „Nieznany region” w pliku o pięciuset wierszach nie mówi nic: nauczyciel nie wie ani której
    rubryki dotyczy, ani co dokładnie w niej stoi. Nazwa kolumny jest tą samą nazwą, którą widzi
    w nagłówku arkusza i w instrukcji nad formularzem, więc poprawka jest jednym spojrzeniem.
    """
    return f"kolumna „{column.label}” („{value}”): {str(detail).strip()}"


#: Który błąd serwisu rejestracji dotyczy której rubryki pliku. Reguły placówki są **te same**, co
#: przy ``/register/`` (``apps.accounts.services._resolve_institution``) i drugiej ich kopii tu nie
#: ma; import dokłada wyłącznie to, czego formularz nie potrzebuje – wskazanie kolumny.
_INSTITUTION_ERROR_COLUMNS: dict[str, Column] = {
    "INSTITUTION_TYPE_REQUIRED": INSTITUTION_TYPE_COLUMN,
    "INSTITUTION_TYPE_NOT_ALLOWED": INSTITUTION_TYPE_COLUMN,
    "COUNTRY_REQUIRED": COUNTRY_COLUMN,
    "COUNTRY_INVALID": COUNTRY_COLUMN,
    "CUSTOM_INSTITUTION_NOT_FOUND": INSTITUTION_ID_COLUMN,
    "SCHOOL_NOT_FOUND": INSTITUTION_ID_COLUMN,
    "SCHOOL_REQUIRED": INSTITUTION_NAME_COLUMN,
    "INSTITUTION_NAME_REQUIRED": INSTITUTION_NAME_COLUMN,
}


def _institution_cell(row: ImportRow, column: Column) -> str:
    """Wartość tej rubryki w tym wierszu – do komunikatu, a nie do rozstrzygnięcia."""
    return {
        INSTITUTION_TYPE_COLUMN: row.institution_type,
        INSTITUTION_ID_COLUMN: row.custom_institution_id,
        INSTITUTION_NAME_COLUMN: row.institution_name,
        COUNTRY_COLUMN: row.country,
    }.get(column, "")


def _custom_institution_id(text: str) -> str:
    """Identyfikator placówki z komórki. Arkusz potrafi podać „12.0” – to nadal jest ten numer."""
    value = (text or "").replace(" ", "")
    return value[:-2] if value.endswith(".0") else value


def _names_own_institution(row: ImportRow) -> bool:
    """Czy wiersz wskazuje **własną** placówkę, czy zostaje przy szkole całego pliku.

    Blok „placówka” czytamy jako całość i tylko wtedy, gdy wiersz naprawdę coś nazywa: wskazuje
    wykaz organizatora (identyfikator), wpisuje nazwę albo deklaruje rodzaj placówki spoza wykazu
    publicznego (``FOREIGN``, ``NONE``, ``OTHER``). Sam „typ placówki” równy rodzajowi z wykazu nie
    zastępuje szkoły z formularza – szkoła jest jedna dla całego pliku i tak ma zostać (nazwa
    wchodzi do progu k-anonimowości w publikacji wyników).
    """
    from .models import DIRECTORY_INSTITUTION_TYPES

    if row.custom_institution_id or row.institution_name:
        return True
    return bool(row.institution_type) and row.institution_type not in DIRECTORY_INSTITUTION_TYPES


def _resolve_row_context(rows: list[ImportRow], competition) -> None:
    """Dopisuje wierszom to, co wynika z **konfiguracji konkursu**: region, kategorię i placówkę.

    Funkcja jest w całości warunkowa i to jest jej główna treść: przy konkursie bez flag etapu 2
    żaden wiersz nie ma wypełnionej ani jednej z tych kolumn (bo :func:`extra_columns` ich nie
    wystawiła), więc kończy się na trzech sprawdzeniach napisów i **nie zadaje ani jednego
    zapytania** (§ 5.6).

    Reguły nie są tu przepisane, tylko wywołane: region rozstrzyga
    ``apps.accounts.services._resolve_region`` (§ 1.4.2), kategorię ``Category.auto_for_grade``
    i kod z wiersza (§ 1.2.4), a placówkę ``_resolve_institution`` (§ 1.3.2–1.3.4) – ta sama
    funkcja i te same komunikaty, co przy ``/register/``. Import dokłada wyłącznie to, czego
    formularz nie potrzebuje: wskazanie rubryki i wartości, bo błąd dotyczy wiersza w arkuszu.

    Wyniki są **na wiersz** i nie jadą w koszyku: zapis liczy je od nowa, tak samo jak decyzję
    „założyć czy dowiązać”.
    """
    if competition is None:
        return
    wants_region = any(row.region_code for row in rows)
    wants_category = competition.has_feature(CATEGORIES_FLAG)
    wants_institution = any(
        row.institution_type or row.custom_institution_id or row.institution_name or row.country
        for row in rows
    )
    if not (wants_region or wants_category or wants_institution):
        return
    from .services import registration_profile

    profile = registration_profile(competition)
    if wants_region:
        _resolve_rows_region(rows, competition, profile)
    if wants_category:
        _resolve_rows_category(rows, competition, profile)
    if wants_institution:
        _resolve_rows_institution(rows, competition, profile)


def _resolve_rows_region(rows: list[ImportRow], competition, profile) -> None:
    """Region wiersza z kolumny „region”: kod podziału konkursu albo nazwa województwa.

    Pamięć podręczna jest na **wartość komórki**, a nie na wiersz: lista klasowa ma zwykle jeden
    region na cały plik, więc pięćset wierszy kosztuje jedno zapytanie. Nierozpoznana wartość
    schodzi do dzisiejszej reguły (lista województw) i dopiero jej odmowa jest błędem wiersza –
    konkurs z włączoną flagą, ale bez własnego podziału, ma zachowywać się jak przed nią.
    """
    from .services import _resolve_region

    resolved: dict[str, tuple[str, object] | DomainError] = {}
    for row in rows:
        if not row.region_code:
            continue
        if row.region_code not in resolved:
            try:
                resolved[row.region_code] = _resolve_region(
                    competition, row.region_code, row.region_code, profile=profile
                )
            except DomainError as exc:
                resolved[row.region_code] = exc
        found = resolved[row.region_code]
        if isinstance(found, DomainError):
            row.errors.append(_column_error(REGION_COLUMN, row.region_code, found.detail))
        else:
            row.district, row.region = found


def _resolve_rows_category(rows: list[ImportRow], competition, profile) -> None:
    """Kategoria startowa wiersza: kod z kolumny „kategoria” albo reguła klas (§ 1.2.4).

    Kod wskazany wprost wygrywa zawsze. Gdy go nie ma, kategorię liczy ``Category.auto_for_grade``
    – ale **tylko** wtedy, gdy konkurs nie pozwala wskazać jej uczestnikowi: w konkursie, w którym
    kategoria jest wyborem zawodnika, wpisanie mu jej z klasy byłoby podjęciem decyzji za niego.
    Zapytania liczą się na **klasy**, a tych jest najwyżej pięć.
    """
    from apps.competitions.models import Category

    available = {
        category.code: category
        for category in Category.objects.for_competition(competition).filter(is_active=True)
    }
    auto: dict[int | None, object] = {}
    for row in rows:
        if row.errors:
            continue
        if row.category_code:
            found = available.get(row.category_code)
            if found is None:
                row.errors.append(
                    _column_error(CATEGORY_COLUMN, row.category_code, "nie ma takiej kategorii.")
                )
            else:
                row.category = found
            continue
        if profile.participant_picks_category:
            continue
        if row.grade not in auto:
            auto[row.grade] = Category.auto_for_grade(competition, row.grade)
        row.category = auto[row.grade]


def _resolve_rows_institution(rows: list[ImportRow], competition, profile) -> None:
    """Placówka wiersza: wykaz organizatora, wolny tekst albo sam kraj (§ 1.3.2–1.3.4).

    Dwie drogi, bo blok „placówka” odpowiada na dwa różne pytania. Wiersz, który **nazywa** swoją
    placówkę, idzie przez ``_resolve_institution`` – czyli dokładnie tam, gdzie idzie rejestracja,
    z tymi samymi odmowami. Wiersz, który podaje wyłącznie rodzaj placówki albo kraj, zostaje przy
    szkole całego pliku, a te dwie wartości są tylko sprawdzane: rodzaj listą dopuszczonych,
    kraj tą samą regułą ISO, co przy ``/register/``.

    Pamięć podręczna jest na **treść bloku**: klasa wyjeżdżająca z jednej uczelni partnerskiej ma
    ten sam blok w każdym wierszu, więc kosztuje jedno zapytanie, a nie pięćset.
    """
    from .services import _require_country, _resolve_institution

    allowed = profile.institution_types()
    cache: dict[tuple[str, str, str, str], dict | DomainError] = {}
    for row in rows:
        key = (row.institution_type, row.custom_institution_id, row.institution_name, row.country)
        if not any(key):
            continue
        if row.institution_type and row.institution_type not in allowed:
            row.errors.append(
                _column_error(
                    INSTITUTION_TYPE_COLUMN,
                    row.institution_type,
                    # Zdanie jest przepisane **znak w znak** z serwisu rejestracji: ta sama odmowa
                    # ma brzmieć tak samo niezależnie od tego, czy przyszła z formularza, czy z pliku.
                    "Ten rodzaj placówki nie jest dopuszczony w tym konkursie.",
                )
            )
            continue
        if not _names_own_institution(row):
            # Sam kraj (albo sam rodzaj z wykazu publicznego): szkoła zostaje ta z formularza,
            # a jedyną rzeczą do rozstrzygnięcia jest zapis kraju – tą samą regułą, co rejestracja.
            try:
                country = _require_country(
                    row.country, profile=profile, institution_type=row.institution_type or allowed[0]
                )
            except DomainError as exc:
                row.errors.append(_column_error(COUNTRY_COLUMN, row.country, exc.detail))
                continue
            if country:
                row.institution = {"country": country}
            continue
        if key not in cache:
            fields = {
                "institution_type": row.institution_type,
                "custom_institution_id": _custom_institution_id(row.custom_institution_id) or None,
                "institution_name": row.institution_name,
                # ``school`` jest tu tą samą komórką, co ``institution_name``: wolny tekst z pliku
                # wchodzi obiema drogami, bo ``_resolve_institution`` czyta ``school`` dla wykazu
                # publicznego, a ``institution_name`` dla placówki spoza wykazu.
                "school": row.institution_name,
                "school_id": None,
                "country": row.country,
            }
            try:
                cache[key] = _resolve_institution(profile, fields, competition=competition)
            except DomainError as exc:
                cache[key] = exc
        found = cache[key]
        if isinstance(found, DomainError):
            column = _INSTITUTION_ERROR_COLUMNS.get(found.machine_code, INSTITUTION_NAME_COLUMN)
            row.errors.append(_column_error(column, _institution_cell(row, column), found.detail))
        else:
            row.institution = dict(found)


# --- rozstrzygnięcie wierszy ------------------------------------------------------------------


def validate_rows(rows: list[ImportRow], *, competition=None) -> list[ImportRow]:
    """Dopisuje każdemu wierszowi rozstrzygnięcie: założyć, dowiązać czy pominąć.

    Trzy reguły, wszystkie wymuszone tutaj, a nie w widoku, bo droga jest jedna dla nauczyciela
    i koordynatora, a podgląd musi pokazywać **dokładnie to**, co zrobi zatwierdzenie:

    1. **adres powtórzony w pliku** liczy się raz. Porównanie bez względu na wielkość liter, bo
       „Jan.Kowalski@…” i „jan.kowalski@…” to jedna skrzynka i jedno konto,
    2. **adres zajęty przez konto uczestnika** nie zakłada drugiego konta – dopisuje opiekuna do
       istniejącego profilu. Uczeń, który zarejestrował się sam tydzień wcześniej, nie może przez
       import stracić swojego konta ani dostać drugiego,
    3. **adres zajęty przez konto bez profilu uczestnika** (recenzent, koordynator, opiekun)
       zostaje pominięty. Dorobienie takiemu koncu profilu uczestnika zmieniałoby komuś rolę
       w zawodach na podstawie pliku wgranego przez osobę trzecią.

    „Konto uczestnika” znaczy od tej zmiany „konto z profilem **w tym konkursie**”: nauczyciel
    importujący klasę do olimpiady A nie może dostać wiersza „dowiązać” dlatego, że uczeń startuje
    w olimpiadzie B – tam jest jego profil, jego zgody i jego opiekun, a tutaj nie ma jeszcze nic.
    ``competition=None`` bierze konkurs z kontekstu (``default_competition``), tak samo jak zapis.

    Przed rozstrzygnięciem idzie :func:`_resolve_row_context`, czyli kolumny dokładane przez
    konkurs (region, kategoria, placówka). Kolejność nie jest dowolna: ich odmowy są **błędami
    wiersza**, a wiersz z błędem ma zostać pominięty i nie ma po co pytać o jego adres.
    """
    from .services import default_competition

    if competition is None:
        competition = default_competition()
    _resolve_row_context(rows, competition)
    emails = [row.email for row in rows if row.email and not row.errors]
    # Jedno zapytanie na cały plik: zbiór adresów, które mają w **tym** konkursie profil
    # uczestnika, i zbiór adresów zajętych w ogóle. Dwa zbiory, bo prowadzą do dwóch różnych
    # rozstrzygnięć (dowiązać / pominąć), a nie do jednego z warunkiem.
    taken = set(User.objects.filter(email__in=emails).values_list("email", flat=True))
    with_profile = set(
        Participant.objects.filter(user__email__in=emails, competition=competition).values_list(
            "user__email", flat=True
        )
    )
    seen: set[str] = set()
    for row in rows:
        if row.errors:
            row.action = ACTION_SKIP
            continue
        if row.email in seen:
            row.action = ACTION_SKIP
            row.errors.append("adres powtórzony w pliku")
            continue
        seen.add(row.email)
        if row.email not in taken:
            row.action = ACTION_CREATE
            if row.birth_year is not None and is_minor(row.birth_year) and not row.guardian_email:
                row.notes.append(
                    "uczeń niepełnoletni bez adresu opiekuna prawnego – zgodę opiekuna uzupełni "
                    "sam po przyjęciu zaproszenia"
                )
            continue
        if row.email not in with_profile:
            row.action = ACTION_SKIP
            row.errors.append("adres należy do konta, które nie jest kontem uczestnika")
            continue
        row.action = ACTION_LINK
        row.notes.append("konto już istnieje – zostanie dopisane do opiekuna, bez zakładania nowego")
    return rows


def pack_rows(rows: list[ImportRow]) -> str:
    """Koszyk wierszy niesiony między podglądem a zatwierdzeniem, podpisany ``django.core.signing``.

    Dlaczego nie sesja: lista klasowa jest zbiorem danych osobowych **cudzych** i nie ma powodu,
    żeby leżała w naszej bazie sesji przez cały czas namysłu nauczyciela. Podpisany koszyk żyje
    w jego przeglądarce, wygasa po dwóch godzinach i nie zostawia po sobie ani jednego wiersza.

    Podpis nie jest szyfrowaniem i nie udaje nim być: treść jest czytelna dla tego, kto ma tę
    stronę otwartą – czyli dla osoby, która przed chwilą sama wgrała ten plik. Chroni przed czym
    innym: przed podmianą listy między podglądem a zatwierdzeniem.
    """
    return signing.dumps([row.payload() for row in rows], salt=PREVIEW_SALT, compress=True)


def unpack_rows(token: str) -> list[ImportRow]:
    """Odczytuje koszyk. Podpis zły albo przeterminowany = podgląd trzeba zrobić od nowa."""
    try:
        payload = signing.loads(token, salt=PREVIEW_SALT, max_age=PREVIEW_MAX_AGE)
    except signing.BadSignature as exc:  # obejmuje ``SignatureExpired``
        raise DomainError(
            "Podgląd importu wygasł. Wgraj plik jeszcze raz.",
            "IMPORT_PREVIEW_EXPIRED",
            status.HTTP_400_BAD_REQUEST,
        ) from exc
    if not isinstance(payload, list) or len(payload) > MAX_ROWS:
        raise DomainError(
            "Podgląd importu jest nieczytelny. Wgraj plik jeszcze raz.",
            "IMPORT_PREVIEW_INVALID",
            status.HTTP_400_BAD_REQUEST,
        )
    return [ImportRow.from_payload(item) for item in payload if isinstance(item, dict)]


def preview_upload(
    upload, *, with_supervisor: bool, default_supervisor_email: str = "", competition=None
) -> ImportPreview:
    """Pełna droga „plik → podgląd”: odczyt, oczyszczenie, rozstrzygnięcie i koszyk.

    Konkurs wchodzi **jednym** argumentem i dwa razy: raz decyduje o zestawie kolumn
    (:func:`extra_columns`), raz o rozstrzygnięciu wierszy. Dwa różne konkursy w tych dwóch
    miejscach znaczyłyby plik, którego nagłówek przyjęliśmy, a treści nie umiemy przypisać.
    """
    table = read_table(upload)
    rows = parse_table(
        table,
        columns=columns_for(with_supervisor=with_supervisor, competition=competition),
        default_supervisor_email=default_supervisor_email,
    )
    validate_rows(rows, competition=competition)
    importable = [row for row in rows if row.action != ACTION_SKIP]
    return ImportPreview(rows=rows, token=pack_rows(importable))


# --- zapis ------------------------------------------------------------------------------------


def _create_invited_user(row: ImportRow) -> User:
    """Konto w stanie „zaproszony”: nieaktywne i **bez używalnego hasła**.

    Hasła nie losujemy „na wszelki wypadek” – ``set_unusable_password`` jest tu dokładniejsze:
    nie istnieje ciąg znaków, którym dałoby się na to konto wejść, więc nie ma czego wycieknąć
    ani czego przypadkiem wysłać w liście. Uczeń ustawia pierwsze hasło sam, pod linkiem
    zaproszenia; ta sama droga, co przy „Nie pamiętasz hasła?”.
    """
    user = User(
        email=row.email,
        first_name=row.first_name,
        last_name=row.last_name,
        is_active=False,
    )
    user.set_unusable_password()
    user.save()
    return user


def _placement(row: ImportRow, competition, *, school_name: str, school_ref, district: str) -> dict:
    """Kolumny profilu opisujące „skąd startuje ten uczeń”: placówka, kraj, okręg i region.

    Wartością domyślną każdej z nich jest **dzisiejsza** wartość: szkoła z formularza, jej
    województwo, brak kraju, brak dowiązania do wykazu organizatora i brak regionu. Wiersz
    podmienia z tego tylko to, co sam nazwał – a Konkurs #1 nie nazywa niczego, więc dostaje
    dokładnie ten sam zapis, co przed etapem 2 (§ 0.1).

    Kolejność źródeł okręgu jest przepisana z § 1.3.4 i ma trzy stopnie, bo tyle jest coraz
    słabszych przesłanek: kolumna „region” (ktoś to napisał wprost), region placówki z wykazu
    organizatora (``CustomInstitution.region_code``) albo „poza Polską” dla placówki zagranicznej,
    i dopiero na końcu województwo szkoły całego pliku.
    """
    fields = {
        "school": school_name,
        "school_ref": school_ref,
        "custom_institution_ref": None,
        "institution_name": "",
        "country": "",
    }
    if row.institution:
        fields.update(row.institution)
    if row.district:
        return {**fields, "district": row.district, "region": row.region}
    fallback = _fallback_district(row, competition, fields, district)
    return {**fields, **_district_and_region(competition, fallback)}


def _fallback_district(row: ImportRow, competition, fields: dict, district: str) -> str:
    """Okręg wiersza, gdy kolumny „region” w pliku nie było – dzisiejsza reguła i jej wyjątek.

    Wiersz, który **nie nazwał** własnej placówki, dostaje województwo szkoły całego pliku: to jest
    dzisiejsza linia i ani jeden znak w niej się nie zmienia.

    Wiersz, który placówkę nazwał, dzisiejszej odpowiedzi mieć nie może – uczeń uczelni partnerskiej
    albo szkoły w Berlinie nie startuje z województwa szkoły nauczyciela. Przy własnym podziale
    (``custom_regions``) odpowiedzią jest region wykazu organizatora albo „poza Polską”; **bez**
    podziału odpowiedzi nie ma i kolumna zostaje pusta. Pusta, a nie zgadnięta: ``district`` czyta
    kilkadziesiąt miejsc jako województwo z zamkniętej listy (filtry panelu, eksporty, reguła
    konfliktu interesów), więc kod spoza tej listy byłby tam wartością, której nikt nie umie
    porównać. O region dopyta uczeń przy przyjęciu zaproszenia – tak samo, jak przy szkole spoza
    rejestru.
    """
    from .regions import ABROAD_CODE
    from .services import CUSTOM_REGIONS_FLAG

    custom = fields.get("custom_institution_ref")
    if custom is None and not _names_own_institution(row):
        return district
    if competition is None or not competition.has_feature(CUSTOM_REGIONS_FLAG):
        return ""
    if custom is not None:
        return (custom.region_code or "").strip()
    return ABROAD_CODE if row.institution_type == "FOREIGN" else ""


def _district_and_region(competition, district: str) -> dict:
    """Para „okręg i region” dla wartości wyliczonej z placówki – jedno miejsce na denormalizację.

    Przy wyłączonej fladze ``custom_regions`` oddaje dzisiejszy napis i ``None``, **bez zapytania**:
    kolumna ``region`` jest wtedy pusta w każdym wierszu, tak samo jak w rejestracji (§ 1.4.2).
    """
    from .models import region_for_district
    from .services import CUSTOM_REGIONS_FLAG

    if competition is None or not competition.has_feature(CUSTOM_REGIONS_FLAG):
        return {"district": district, "region": None}
    found = region_for_district(competition, district)
    if found is None:
        return {"district": district, "region": None}
    return {"district": found.code, "region": found}


def _assign_categories(stage, pairs: list[tuple[int, object]]) -> None:
    """Kategoria startowa na **istniejących** wpisach do etapu (§ 1.2.4).

    ``update`` na zawężonym zbiorze, a nie ``get_or_create``: import nikogo do etapu nie zapisuje
    (patrz :func:`import_students`), więc wiersz, którego nie ma, ma **nie powstać**. Zapytań jest
    tyle, ile różnych kategorii w pliku – czyli najwyżej tyle, ile konkurs ich ma.
    """
    from apps.competitions.models import StageEntry

    by_category: dict[int, list[int]] = {}
    for participant_id, category in pairs:
        if category is not None:
            by_category.setdefault(category.pk, []).append(participant_id)
    for category_id, participant_ids in by_category.items():
        StageEntry.objects.filter(stage=stage, participant_id__in=participant_ids).update(
            category_id=category_id
        )


@transaction.atomic
def import_students(
    rows: list[ImportRow],
    *,
    school_name: str,
    school_ref=None,
    default_supervisor_email: str = "",
    stage=None,
    actor=None,
    request=None,
) -> dict:
    """Zakłada konta zaproszonych i dowiązuje istniejące. Zwraca licznik rozstrzygnięć.

    Szkoła jest jedna dla całego pliku i tak ma być: nauczyciel importuje **swoją** klasę, a nazwa
    szkoły wchodzi do grupowania w publikowanych wynikach (próg k-anonimowości), więc musi być
    zapisana identycznie u wszystkich uczniów placówki. Kolumny „szkoła” w pliku nie ma i nie
    będzie – dwadzieścia ręcznie wpisanych wariantów tej samej nazwy to dwadzieścia szkół
    w statystyce. Kolumna **„placówka”** (etap 2, § 1.3.2) jest czym innym i dlatego wolno jej
    istnieć: nie jest drugim zapisem tej samej szkoły, tylko wskazaniem, że **ten** uczeń startuje
    z innej placówki niż reszta listy – z wykazu organizatora albo spoza Polski.

    Województwo bierzemy ze **szkoły z rejestru**, gdy taka jest. Gdy nie ma (szkoła wpisana
    ręcznie), zostaje puste i pyta o nie uczeń przy przyjęciu zaproszenia – nauczyciel i tak nie
    odpowiada za to pole, a zgadywanie wstawiłoby do bazy wartość, której nikt nie potwierdził.

    ``stage`` jest nieobowiązkowy i **nikogo do etapu nie zapisuje**: import nie tworzy wpisów
    i tworzyć ich nie będzie (konto bez zgód nie ma prawa startować, patrz dokumentacja modułu).
    Podany etap znaczy wyłącznie „uzupełnij kategorię na wpisach, które już istnieją” (§ 1.2.4) –
    czyli dokładnie to, po co koordynator wgrywa listę z kolumną „kategoria” do trwającej edycji.
    """
    from .services import create_participant_with_public_code, default_competition, grant_role
    from .supervisors import set_supervisor_email

    # Konkurs importu ustalamy **raz** dla całego pliku, i to **przed** rozstrzygnięciem wierszy:
    # jedna lista klasowa nie ma prawa rozsypać się po dwóch konkursach, a podgląd i zapis mają
    # pytać o istniejące profile w tym samym konkursie.
    competition = default_competition()
    validate_rows(rows, competition=competition)
    district = getattr(school_ref, "voivodeship", "") or ""
    now = timezone.now()
    created = 0
    linked = 0
    categorised: list[tuple[int, object]] = []
    for row in rows:
        if row.action == ACTION_CREATE:
            user = _create_invited_user(row)
            grant_role(user, GROUP_PARTICIPANT, competition=competition)
            placement = _placement(
                row, competition, school_name=school_name, school_ref=school_ref, district=district
            )
            participant = create_participant_with_public_code(
                user=user,
                competition=competition,
                grade=row.grade,
                birth_year=row.birth_year,
                phone=row.phone,
                guardian_email=row.guardian_email,
                supervisor_email=row.supervisor_email or default_supervisor_email,
                invited_at=now,
                **placement,
            )
            categorised.append((participant.pk, row.category))
            send_invitation(participant, request=request)
            # Wpis per konto, bo to jest zdarzenie dotyczące **tej** osoby: bez niego koordynator
            # patrzący na konto nie wie, skąd się wzięło. W ``diff`` nie ma ani adresu, ani
            # nazwiska – kogo dotyczy, mówi ``target_id``.
            audit(actor, "participant.invited_by_import", participant, {"row": row.number}, request=request)
            created += 1
        elif row.action == ACTION_LINK:
            # Profil **tego** konkursu: po zmianie z wydania D jedno konto ma tyle profili,
            # w ilu konkursach startuje, więc ``get`` bez zawężenia podniósłby
            # ``MultipleObjectsReturned`` na uczniu dwóch olimpiad.
            participant = Participant.objects.select_related("user").get(
                user__email=row.email, competition=competition
            )
            set_supervisor_email(
                participant,
                row.supervisor_email or default_supervisor_email,
                actor=actor,
                request=request,
            )
            # Profilu **nie** nadpisujemy: uczeń, który zarejestrował się sam, podał swoją szkołę,
            # swój region i swoją placówkę, a plik nauczyciela nie jest powodem, żeby mu je zmienić
            # (§ 0.1). Dowiązanie dopisuje opiekuna szkolnego i tyle – tak samo, jak przed etapem 2.
            categorised.append((participant.pk, row.category))
            linked += 1
    if stage is not None:
        _assign_categories(stage, categorised)
    skipped = sum(1 for row in rows if row.action == ACTION_SKIP)
    # Drugi wpis opisuje **przebieg**, a nie konto: to on odpowiada na pytanie „kto i kiedy wgrał
    # listę”. Same liczby, zero adresów – audyt czytają także osoby bez wglądu w listy klasowe.
    audit(
        actor,
        "accounts.students_imported",
        actor if actor is not None else User(pk=0),
        {"created": created, "linked": linked, "skipped": skipped},
        request=request,
    )
    logger.info("Import uczniów: założono %s, dowiązano %s, pominięto %s.", created, linked, skipped)
    return {"created": created, "linked": linked, "skipped": skipped}


# --- zaproszenie ------------------------------------------------------------------------------


def make_invite_token(participant: Participant) -> str:
    """Token wiązany z parą (konto, adres e-mail) – tak samo, jak token aktywacyjny.

    Wiązanie z adresem jest tu regułą bezpieczeństwa, a nie ozdobą: gdyby koordynator poprawił
    adres konta (literówka w liście klasowej), list wysłany pod stary adres przestaje otwierać
    konto, które należy już do kogoś innego.
    """
    return signing.dumps({"pk": participant.user_id, "email": participant.user.email}, salt=INVITE_SALT)


def _invalid_invite() -> DomainError:
    """Jeden komunikat na każdy powód odrzucenia – bez wskazywania, który to był."""
    return DomainError(INVALID_TOKEN_MESSAGE, "INVITE_TOKEN_INVALID", status.HTTP_400_BAD_REQUEST)


def read_invite_token(token: str) -> Participant:
    """Uczestnik wskazany zaproszeniem. ``DomainError`` przy tokenie złym, wygasłym albo zużytym.

    Zużyty znaczy „konto jest już aktywne”: token nie ma stanu w bazie i nie musi mieć –
    jednorazowość bierze się z tego, że drugie wejście nie ma czego zmienić. Osobnego komunikatu
    dla tego przypadku nie ma świadomie, bo rozróżnienie „wygasł” od „już użyty” mówiłoby
    obcemu, że konto pod tym adresem istnieje i działa.
    """
    try:
        payload = signing.loads(token, salt=INVITE_SALT, max_age=INVITE_MAX_AGE)
    except signing.BadSignature as exc:  # obejmuje ``SignatureExpired``
        raise _invalid_invite() from exc
    if not isinstance(payload, dict) or not payload.get("pk"):
        raise _invalid_invite()
    participant = (
        Participant.objects.select_related("user", "school_ref")
        .filter(user_id=payload["pk"], invited_at__isnull=False)
        .first()
    )
    if participant is None:
        raise _invalid_invite()
    user = participant.user
    if user.email != (payload.get("email") or "").strip().lower():
        raise _invalid_invite()
    if user.email_verified_at is not None or user.is_active:
        raise _invalid_invite()
    return participant


def send_invitation(participant: Participant, *, request=None) -> None:
    """Kolejkuje list z zaproszeniem i zapisuje, kiedy poszedł.

    Treść jest w szablonach (``templates/registration/invite_student_*.txt``), a nie w tym pliku,
    bo list czyta uczeń – to jest tekst interfejsu, a nie komunikat serwisu. Poza imieniem
    i nazwą szkoły nie ma w nim danych osobowych: skrzynka bywa cudza, bo adres przepisał ktoś
    inny z listy klasowej.
    """
    link = absolute_url(reverse("web:student-invite", args=[make_invite_token(participant)]), request)
    context = {
        "first_name": participant.user.first_name,
        "school": participant.school,
        "link": link,
        "days": INVITE_DAYS,
    }
    subject = render_to_string(INVITE_SUBJECT_TEMPLATE, context).strip().replace("\n", " ")
    queue_mail(subject, render_to_string(INVITE_BODY_TEMPLATE, context), participant.user.email)
    Participant.objects.filter(pk=participant.pk).update(invitation_sent_at=timezone.now())
    participant.invitation_sent_at = timezone.now()


def resend_invitation(participant: Participant, *, actor=None, request=None) -> Participant:
    """Wysyła zaproszenie ponownie. Odmowa dla konta, które zaproszenia nie ma albo już działa.

    Odmowa jest jawna, a nie cicha: nauczyciel klika „wyślij ponownie” właśnie dlatego, że nie
    wie, co się dzieje, i „nic się nie stało” byłoby najgorszą z możliwych odpowiedzi.
    """
    if participant.invited_at is None:
        raise DomainError(
            "To konto nie powstało z importu – nie ma zaproszenia do wysłania.",
            "INVITE_NOT_IMPORTED",
            status.HTTP_409_CONFLICT,
        )
    if participant.user.email_verified_at is not None or participant.user.is_active:
        raise DomainError(
            "Ten uczeń przyjął już zaproszenie – konto jest aktywne.",
            "INVITE_ALREADY_ACCEPTED",
            status.HTTP_409_CONFLICT,
        )
    send_invitation(participant, request=request)
    audit(actor, "participant.invitation_resent", participant, {}, request=request)
    return participant


#: Stany widoczne w panelu opiekuna i na liście kont. Wartości są kluczami, etykietę składa
#: szablon – ten moduł nie zna języka interfejsu (ta sama zasada, co w ``apps.accounts.guardian``).
STATE_INVITED = "invited"
STATE_ACTIVE = "active"
STATE_SELF_REGISTERED = "self"


def invitation_state(participant: Participant) -> str:
    """Stan zaproszenia jednego ucznia – trzy różne rzeczy, nie dwie.

    „Zaproszony” i „aktywny” dotyczą kont z importu; konto, które uczeń założył sam, nie jest
    „aktywnym zaproszeniem”, tylko czymś innym i panel ma je tak podpisać. Sklejenie obu kazałoby
    nauczycielowi zgadywać, komu wypada przypomnieć o liście, a komu nie ma o czym.
    """
    if participant.invited_at is None:
        return STATE_SELF_REGISTERED
    if participant.user.email_verified_at is None:
        return STATE_INVITED
    return STATE_ACTIVE


@transaction.atomic
def accept_invitation(
    participant: Participant,
    *,
    password: str,
    phone: str,
    district: str,
    given: dict[str, bool],
    request=None,
) -> Participant:
    """Przyjęcie zaproszenia: hasło ucznia, brakujące dane, komplet zgód i aktywacja konta.

    Kolejność nie jest dowolna. Zgody sprawdzamy **przed** zapisaniem czegokolwiek (tak samo, jak
    w ``register_participant``), bo konto bez kompletu zgód nie ma prawa stać się aktywne nawet
    na chwilę wewnątrz transakcji. Aktywacja idzie na końcu i tą samą funkcją, co kliknięcie
    linku z listu (``mark_activated``) – dzięki temu potwierdzenie adresu trafia także do tabeli
    allauth i uczeń nie traci świeżo ustawionego hasła przy pierwszym logowaniu Google'em.

    O co pytamy ucznia, skoro nauczyciel podał już dane: o **hasło** (poświadczenie ma być jego),
    o **województwo** (nie ma go w liście klasowej, a przy szkole spoza rejestru nie ma skąd go
    wziąć) i o **telefon** (nauczyciel bywa go nie znać, a organizator dzwoni w dniu zawodów).
    Imienia, nazwiska, szkoły, klasy i rocznika nie pytamy drugi raz – to są dane, które
    nauczyciel zna z dziennika i których uczeń nie ma po co przepisywać.

    ``given`` przychodzi w postaci **pól formularza** (``terms_consent``, ``gdpr_consent``, …),
    dokładnie tak, jak w ``register_participant`` – kontrakt jest wspólny, żeby widok nie musiał
    wiedzieć, jak nazywają się rodzaje zgód w modelu dowodowym.
    """
    from .consents import ConsentSource, given_from_fields
    from .services import _require_voivodeship, _validate_password_or_raise, record_consents

    user = participant.user
    # ``given`` przychodzi z formularza, czyli z nazwami **pól** (``terms_consent`` …), a serwis
    # zgód mówi **rodzajami** (``TERMS`` …). Mapuje je jedna funkcja, ta sama, co w rejestracji –
    # dzięki niej nazwa pola i rodzaj zgody nie mają jak się rozjechać.
    given = given_from_fields(given)
    validate_consents_for(participant, given)
    district = _require_voivodeship(district, required=True)
    phone = normalize_phone(phone)
    _validate_password_or_raise(password, user)

    participant.district = district
    participant.phone = phone
    participant.save(update_fields=["district", "phone"])
    user.set_password(password)
    user.save(update_fields=["password"])
    record_consents(participant, given, source=ConsentSource.WEB, request=request)
    mark_activated(user, actor=user, action="account.invitation_accepted", request=request)
    participant.refresh_from_db()
    return participant


def validate_consents_for(participant: Participant, given: dict[str, bool]) -> None:
    """Komplet zgód wymaganych od ucznia o tym roczniku – ta sama reguła, co w rejestracji.

    Osobna funkcja wyłącznie dla czytelności wywołania: reguła mieszka w
    ``apps.accounts.services.validate_consents`` i nie jest tu powtórzona.
    """
    from .services import validate_consents

    validate_consents(given, birth_year=participant.birth_year)
