"""Raport obecności Google Meet: odczyt pliku (CSV/XLSX), pobranie arkusza z linku, czasy trwania.

Moduł odpowiada na jedno pytanie: **kto i jak długo był na spotkaniu** według raportu, który Meet
wysyła organizatorowi po zajęciach (funkcja „Śledzenie obecności”, opis formatu: WS-ATT-02 § 1).
Nie wie nic o uczestnikach olimpiady ani o warsztatach – dopasowanie osób do kont i zapis obecności
robi ``apps.cms.meet_import``. Dzięki temu parser testuje się bez bazy, a jedyną drogą danych
z raportu do bazy jest wynik dopasowania (identyfikatory uczestników).

Dlaczego parser jest tolerancyjny, a nie przypięty do jednego układu kolumn. Google nie publikuje
formatu raportu: kolumny zależą od języka interfejsu organizatora, kolejność bywa różna, nad tabelą
stoi czasem blok z kodem i datą spotkania, a koordynator często dostaje arkusz już przerobiony
(sklejony z kilku kart, wyeksportowany przez dodatek, zapisany w Excelu). Odmowa każdego wariantu,
którego nie przewidzieliśmy, oznaczałaby powrót do ręcznego przepisywania kratek.

Dane osobowe z raportu (e-maile, nazwiska) żyją wyłącznie w obiektach tego modułu w czasie jednego
żądania – nic stąd nie trafia do logów ani do cache'u (WS-ATT-02 § 6).
"""

from __future__ import annotations

import csv
import io
import re
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from datetime import time as dtime
from urllib.parse import parse_qs, urljoin, urlsplit

#: Ile wierszy danych przyjmujemy z jednego raportu. Warsztat online olimpiady to kilkadziesiąt do
#: kilkuset osób; dwa tysiące mieszczą z zapasem każde spotkanie, a odcinają plik, który nie jest
#: raportem jednego spotkania (cała baza adresowa, eksport z innego systemu).
MAX_REPORT_ROWS = 2000

#: Rozmiar pliku i pobrania z linku. Raport dwóch tysięcy osób w CSV to ~200 kB, a w XLSX mniej;
#: pięć megabajtów zostawia zapas na formatowanie i nie pozwala trzymać w pamięci procesu ``web``
#: niczego, co raportem nie jest.
MAX_REPORT_BYTES = 5 * 1024 * 1024

#: W ilu pierwszych wierszach szukamy nagłówka tabeli. Nad tabelą bywa blok „meta” (kod
#: spotkania, tytuł, data, organizator) – kilka wierszy, nie kilkadziesiąt.
HEADER_SCAN_ROWS = 30

#: Czasy połączenia z Google (połączenie, odczyt jednego kawałka) i całe pobranie. Osobny limit
#: całości, bo ``timeout`` biblioteki ``requests`` liczy się **na kawałek** – serwer sączący bajt
#: na sekundę trzymałby proces ``web`` w nieskończoność.
CONNECT_TIMEOUT = 5
READ_TIMEOUT = 10
TOTAL_TIMEOUT = 20

#: Najwięcej przekierowań przy eksporcie. Google oddaje eksport jednym przekierowaniem na
#: ``*.googleusercontent.com``; trzy to zapas, a nie zaproszenie do pętli.
MAX_REDIRECTS = 3

SHEETS_HOST = "docs.google.com"
#: Host, na który Google przekierowuje pobranie eksportu. Wyłącznie poddomeny – sama domena
#: ``googleusercontent.com`` nie serwuje eksportów.
EXPORT_HOST_SUFFIX = ".googleusercontent.com"
#: Hosty strony logowania Google – przekierowanie tam znaczy „arkusz nie jest udostępniony”.
LOGIN_HOSTS = frozenset({"accounts.google.com", "accounts.youtube.com"})

#: Identyfikator arkusza w adresie: ``/spreadsheets/d/<ID>/…`` (także ``/spreadsheets/u/0/d/<ID>``).
#: Długość od 25 znaków odcina link publikacji (``/spreadsheets/d/e/2PACX-…``), w którym „ID”
#: to litera ``e``, i śmieci; prawdziwe ID ma 44 znaki.
SHEET_PATH = re.compile(r"^/spreadsheets/(?:u/\d+/)?d/(?P<id>[A-Za-z0-9_-]{25,100})(?:/.*)?$")
PUBLISHED_PATH = re.compile(r"^/spreadsheets/(?:u/\d+/)?d/e/")
GID = re.compile(r"^\d{1,12}$")

NOT_SHARED_MESSAGE = (
    "Arkusz nie jest udostępniony do odczytu. W Arkuszach Google kliknij „Udostępnij” i ustaw "
    "„Każdy, kto ma link – może wyświetlać” albo pobierz raport jako CSV (Plik → Pobierz → "
    "Wartości rozdzielone przecinkami) i wgraj plik."
)
DOWNLOAD_INSTEAD = "Pobierz raport jako CSV (Plik → Pobierz → Wartości rozdzielone przecinkami) i wgraj plik."


class MeetReportError(Exception):
    """Odmowa z komunikatem dla koordynatora – raport nieczytelny, link niedozwolony, pobranie nieudane."""

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


# --- nagłówki ------------------------------------------------------------------------------------


def _header_key(text) -> str:
    """Postać porównawcza nagłówka: małe litery, bez dopisku w nawiasie, dwukropka i podwójnych spacji."""
    value = unicodedata.normalize("NFKC", str(text or "")).casefold()
    value = re.sub(r"\(.*?\)|\[.*?\]", " ", value)
    value = value.replace("_", " ").replace(":", " ").replace("*", " ")
    return " ".join(value.split())


#: Warianty nagłówków raportu w interfejsie polskim i angielskim (WS-ATT-02 § 1.1).
HEADERS: dict[str, frozenset[str]] = {
    "first": frozenset({"imię", "imie", "first name", "firstname", "given name"}),
    "last": frozenset({"nazwisko", "last name", "lastname", "surname", "family name"}),
    "full": frozenset(
        {
            "imię i nazwisko",
            "imie i nazwisko",
            "uczestnik",
            "nazwa",
            "nazwa wyświetlana",
            "nazwa uczestnika",
            "name",
            "full name",
            "participant",
            "participant name",
            "display name",
        }
    ),
    "email": frozenset(
        {
            "adres e-mail",
            "adres email",
            "adresy e-mail",
            "e-mail",
            "e-maile",
            "email",
            "emails",
            "email address",
            "e-mail address",
            "mail",
            "participant email",
        }
    ),
    "duration": frozenset(
        {
            "czas trwania",
            "łączny czas",
            "laczny czas",
            "czas obecności",
            "czas obecnosci",
            "czas w spotkaniu",
            "duration",
            "total duration",
            "time in call",
            "attended duration",
            "attendance duration",
        }
    ),
    "joined": frozenset(
        {
            "godzina dołączenia",
            "czas dołączenia",
            "dołączenie",
            "dołączono",
            "dołączył",
            "time joined",
            "join time",
            "joined",
            "joined at",
            "first seen",
        }
    ),
    "exited": frozenset(
        {
            "godzina wyjścia",
            "godzina opuszczenia",
            "czas opuszczenia",
            "czas wyjścia",
            "wyjście",
            "opuścił",
            "time exited",
            "time left",
            "leave time",
            "exit time",
            "left",
            "left at",
            "last seen",
        }
    ),
}


def map_header(row: list) -> dict[str, int] | None:
    """Wiersz → ``{rodzaj kolumny: numer}``, o ile to jest nagłówek raportu; inaczej ``None``.

    Nagłówek musi mieć kolumnę **osoby** (e-mail, imię, nazwisko albo imię i nazwisko) i kolumnę
    **czasu** (czas trwania albo dołączenie). Jedna z dwóch to za mało: blok meta („Kod spotkania:
    abc-defg-hij”, „Data: 12.11.2026”) potrafi mieć w komórce słowo „nazwa” albo „czas trwania”.
    Pierwsza pasująca kolumna wygrywa – raport z dwiema kolumnami „e-mail” jest raportem przerobionym
    i jego pierwsza kolumna jest tą, którą Meet wypełnił.
    """
    mapping: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _header_key(cell)
        if not key:
            continue
        for kind, variants in HEADERS.items():
            if key in variants and kind not in mapping:
                mapping[kind] = index
                break
    has_person = bool({"email", "first", "last", "full"} & mapping.keys())
    has_time = bool({"duration", "joined"} & mapping.keys())
    return mapping if has_person and has_time else None


# --- czasy ---------------------------------------------------------------------------------------

#: Jednostki w zapisie słownym: polski i angielski interfejs, skróty i pełne słowa.
_UNIT_SECONDS = (
    (re.compile(r"^(?:godz|godzin[ay]?|godzina|h|hr|hrs|hour|hours|g)$"), 3600),
    (re.compile(r"^(?:min|mins|minut[ay]?|minuta|minute|minutes|m)$"), 60),
    (re.compile(r"^(?:s|sek|sekund[ay]?|sekunda|sec|secs|second|seconds)$"), 1),
)
_WORD_PART = re.compile(r"(\d+(?:[.,]\d+)?)\s*([^\W\d_]+)\.?", re.UNICODE)
_CLOCK = re.compile(r"^(\d{1,3}):(\d{2})(?::(\d{2}))?$")
_NUMBER = re.compile(r"^\d+(?:[.,]\d+)?$")


def parse_duration(value) -> int | None:
    """Czas trwania → sekundy albo ``None`` (pusty albo nieczytelny).

    Postacie z WS-ATT-02 § 1.2: „1 godz. 23 min”, „1 hr 5 min”, „45 min”, „01:05:00”, „1:05”
    (godziny:minuty – raport nie podaje sekund bez godzin), sama liczba = minuty. Z XLSX przychodzą
    też typy openpyxl: ``timedelta`` (format czasu trwania), ``time`` (komórka czasu) i liczba –
    ułamek doby, gdy mniejsza od jednego (tak Arkusze i Excel trzymają czas trwania), inaczej minuty.
    """
    if value is None or value == "":
        return None
    if isinstance(value, timedelta):
        return max(0, int(value.total_seconds()))
    if isinstance(value, datetime):
        # Excel bez formatu czasu trwania potrafi oddać czas jako „1899-12-30 01:05:00”.
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, dtime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if value < 0:
            return None
        if isinstance(value, float) and value < 1:
            return round(value * 86400)
        return round(value * 60)
    text = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    if not text:
        return None
    clock = _CLOCK.match(text)
    if clock:
        first, second, third = clock.groups()
        if third is not None:
            return int(first) * 3600 + int(second) * 60 + int(third)
        return int(first) * 3600 + int(second) * 60
    if _NUMBER.match(text):
        return round(float(text.replace(",", ".")) * 60)
    total = 0.0
    found = False
    for number, unit in _WORD_PART.findall(text):
        for pattern, seconds in _UNIT_SECONDS:
            if pattern.match(unit):
                total += float(number.replace(",", ".")) * seconds
                found = True
                break
    return round(total) if found else None


_TIME_OF_DAY = re.compile(r"(\d{1,2}):(\d{2})(?::(\d{2}))?\s*(am|pm|a\.m\.|p\.m\.)?", re.IGNORECASE)


def parse_time_of_day(value) -> int | None:
    """Godzina dołączenia/wyjścia → sekundy od północy (data w komórce jest pomijana)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, dtime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Ułamek doby (XLSX) – część po przecinku, bo całość bywa numerem dnia.
        return round((float(value) % 1) * 86400)
    match = _TIME_OF_DAY.search(str(value))
    if not match:
        return None
    hours, minutes, seconds, meridiem = match.groups()
    hours = int(hours)
    if meridiem:
        is_pm = meridiem.lower().startswith("p")
        hours = hours % 12 + (12 if is_pm else 0)
    if hours > 23 or int(minutes) > 59:
        return None
    return hours * 3600 + int(minutes) * 60 + int(seconds or 0)


def span_seconds(joined, exited) -> int | None:
    """„Wyjście − dołączenie” dla raportu bez kolumny czasu trwania; przez północ = +24 h."""
    start, end = parse_time_of_day(joined), parse_time_of_day(exited)
    if start is None or end is None:
        return None
    if end < start:
        end += 86400
    return end - start


def format_duration(seconds: int | None) -> str:
    """Sekundy → „1 godz. 5 min” – tak, jak liczy człowiek, a nie arkusz."""
    if seconds is None:
        return "—"
    minutes = int(seconds) // 60
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours} godz. {minutes} min"
    return f"{minutes} min"


# --- daty spotkania ------------------------------------------------------------------------------

_MONTHS = {
    # polskie – dopełniacz, mianownik i skróty
    "sty": 1,
    "lut": 2,
    "mar": 3,
    "kwi": 4,
    "maj": 5,
    "cze": 6,
    "lip": 7,
    "sie": 8,
    "wrz": 9,
    "paź": 10,
    "paz": 10,
    "lis": 11,
    "gru": 12,
    # angielskie – trzy pierwsze litery
    "jan": 1,
    "feb": 2,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}
_ISO_DATE = re.compile(r"(?<!\d)(\d{4})-(\d{1,2})-(\d{1,2})(?!\d)")
_DMY_DATE = re.compile(r"(?<!\d)(\d{1,2})[./](\d{1,2})[./](\d{4})(?!\d)")
_DAY_MONTH_YEAR = re.compile(r"(?<!\d)(\d{1,2})\.?\s+([^\W\d_]{3,})\.?,?\s+(\d{4})(?!\d)", re.UNICODE)
_MONTH_DAY_YEAR = re.compile(r"\b([^\W\d_]{3,})\.?\s+(\d{1,2}),?\s+(\d{4})(?!\d)", re.UNICODE)


def _month(word: str) -> int | None:
    return _MONTHS.get(word.casefold()[:3])


def _safe_date(year, month, day) -> date | None:
    try:
        return date(int(year), int(month), int(day))
    except TypeError, ValueError:
        return None


def find_date(value) -> date | None:
    """Pierwsza data w tekście komórki albo nazwie pliku; ``None``, gdy nie ma żadnej.

    Postacie z WS-ATT-02 § 3.2. Zapis amerykański „11/12/2026” czytamy jako dzień/miesiąc – raport
    polskiego organizatora; data służy wyłącznie **podpowiedzi** warsztatu, którą koordynator widzi.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = unicodedata.normalize("NFKC", str(value or ""))
    if not text.strip():
        return None
    match = _ISO_DATE.search(text)
    if match and (found := _safe_date(*match.groups())):
        return found
    match = _DMY_DATE.search(text)
    if match:
        day, month, year = match.groups()
        if found := _safe_date(year, month, day):
            return found
    match = _DAY_MONTH_YEAR.search(text)
    if match and (month := _month(match.group(2))):
        if found := _safe_date(match.group(3), month, match.group(1)):
            return found
    match = _MONTH_DAY_YEAR.search(text)
    if match and (month := _month(match.group(1))):
        if found := _safe_date(match.group(3), month, match.group(2)):
            return found
    return None


# --- wynik ---------------------------------------------------------------------------------------


@dataclass
class ReportPerson:
    """Jedna osoba z raportu po zsumowaniu jej wierszy."""

    first_name: str = ""
    last_name: str = ""
    full_name: str = ""
    email: str = ""
    seconds: int = 0
    rows: int = 0
    #: Czy którykolwiek wiersz miał nieczytelny czas – podgląd mówi to wprost zamiast „0 min”.
    unreadable_time: bool = False

    @property
    def label(self) -> str:
        name = " ".join(part for part in (self.first_name, self.last_name) if part) or self.full_name
        return name or self.email or "(bez nazwy)"


@dataclass
class MeetReport:
    people: list[ReportPerson] = field(default_factory=list)
    rows: int = 0
    meeting_date: date | None = None

    @property
    def longest_seconds(self) -> int:
        return max((person.seconds for person in self.people), default=0)


# --- czytanie tabeli -----------------------------------------------------------------------------


def _cell(row: list, mapping: dict[str, int], kind: str):
    index = mapping.get(kind)
    if index is None or index >= len(row):
        return None
    return row[index]


def _text(value) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).split())


#: Domeny Gmaila: kropki w części lokalnej nic nie znaczą, ``+tag`` jest aliasem tej samej skrzynki.
GMAIL_DOMAINS = frozenset({"gmail.com", "googlemail.com"})


def normalize_email(value: str) -> str:
    """Postać porównawcza adresu (WS-ATT-02 § 4.1) – ta sama po stronie raportu i kont.

    Kropki i ``+tag`` zdejmujemy **wyłącznie** w Gmailu: tam z definicji wskazują tę samą skrzynkę
    (uczestnik zapisał się jako ``jan.kowalski+olimpiada@gmail.com``, a Meet pokazuje adres konta
    Google ``jankowalski@gmail.com``). W innych domenach kropka bywa częścią adresu, a plus –
    zwykłym znakiem; zrównanie dwóch różnych skrzynek dałoby obecność cudzemu uczestnikowi.
    """
    text = unicodedata.normalize("NFKC", str(value or "")).strip().casefold()
    local, at, domain = text.rpartition("@")
    if not at:
        return text
    if domain in GMAIL_DOMAINS:
        local = local.split("+", 1)[0].replace(".", "")
        domain = "gmail.com"
    return f"{local}@{domain}"


def name_key(first: str = "", last: str = "", full: str = "") -> str:
    """Imię i nazwisko do porównania: ``normalize_text`` z kont i **posortowane słowa**.

    Sortowanie, bo Meet pokazuje nazwę z konta Google, a tam kolejność bywa odwrócona
    („Kowalski Jan”); dwie różne osoby, których imię i nazwisko są swoją zamianą, to przypadek,
    który i tak trafi do „niejednoznacznych”, jeśli obie są w konkursie.
    """
    from apps.accounts.duplicates import normalize_text

    name = normalize_text(" ".join(part for part in (first, last) if part) or full)
    return " ".join(sorted(name.split()))


def person_key(email: str, first: str, last: str, full: str) -> str:
    """Klucz sumowania wierszy jednej osoby: e-mail, a bez niego imię i nazwisko."""
    if email:
        return "e:" + normalize_email(email)
    name = name_key(first, last, full)
    return "n:" + name if name else ""


def parse_table(table: list[list], filename: str = "") -> MeetReport:
    """Tabela komórek → raport. ``MeetReportError``, gdy nie ma nagłówka albo danych."""
    header_index = None
    mapping: dict[str, int] | None = None
    for index, row in enumerate(table[:HEADER_SCAN_ROWS]):
        mapping = map_header(row)
        if mapping is not None:
            header_index = index
            break
    if header_index is None or mapping is None:
        raise MeetReportError(
            "Nie rozpoznano raportu obecności: brak wiersza nagłówka z kolumnami osoby (imię, nazwisko "
            "albo e-mail) i czasu (czas trwania albo godzina dołączenia). Wgraj arkusz z Google Meet "
            "bez zmian albo dopisz te nagłówki."
        )
    meeting_date = None
    for row in table[:header_index]:
        for cell in row:
            if meeting_date is None:
                meeting_date = find_date(cell)
    people: dict[str, ReportPerson] = {}
    rows = 0
    for row in table[header_index + 1 :]:
        if not any(_text(cell) for cell in row):
            continue
        first = _text(_cell(row, mapping, "first"))
        last = _text(_cell(row, mapping, "last"))
        full = _text(_cell(row, mapping, "full"))
        email = _text(_cell(row, mapping, "email"))
        if email and "@" not in email:
            email = ""
        key = person_key(email, first, last, full)
        if not key:
            # Wiersz bez osoby: podsumowanie pod tabelą („Liczba uczestników: 42”), pusta linia
            # w środku, wiersz numeru telefonu bez nazwy. Nie liczymy go do limitu.
            continue
        rows += 1
        if rows > MAX_REPORT_ROWS:
            raise MeetReportError(
                f"Raport ma więcej niż {MAX_REPORT_ROWS} wierszy – "
                "to nie wygląda na raport jednego spotkania."
            )
        seconds = parse_duration(_cell(row, mapping, "duration")) if "duration" in mapping else None
        if seconds is None:
            # Bez kolumny czasu trwania (albo z pustą komórką) – różnica godzin wyjścia i dołączenia.
            seconds = span_seconds(_cell(row, mapping, "joined"), _cell(row, mapping, "exited"))
        if meeting_date is None:
            meeting_date = find_date(_cell(row, mapping, "joined"))
        person = people.get(key)
        if person is None:
            person = people[key] = ReportPerson(first_name=first, last_name=last, full_name=full, email=email)
        person.rows += 1
        if seconds is None:
            person.unreadable_time = True
        else:
            person.seconds += seconds
    if not people:
        raise MeetReportError("Raport nie ma ani jednego wiersza z osobą pod nagłówkiem.")
    if meeting_date is None and filename:
        meeting_date = find_date(filename)
    return MeetReport(people=list(people.values()), rows=rows, meeting_date=meeting_date)


def decode_csv(data: bytes) -> str:
    """Tekst CSV: BOM UTF-8/UTF-16, UTF-8, UTF-16 bez BOM, inaczej Windows-1250 (WS-ATT-02 § 1.4).

    UTF-16 bez BOM rozpoznajemy po bajtach zerowych: tekst łaciński w UTF-16 ma co drugi bajt zerowy,
    a w żadnym z pozostałych kodowań bajt zerowy w tekście nie występuje. Windows-1250 jest ostatnie,
    bo dekoduje **każdy** ciąg bajtów – to domyślne kodowanie polskiego Excela przy „CSV”.
    """
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    head = data[:2048]
    if head and head.count(b"\x00") * 4 >= len(head):
        even_zeros = head[0::2].count(0)
        odd_zeros = head[1::2].count(0)
        return data.decode("utf-16-be" if even_zeros > odd_zeros else "utf-16-le", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1250", errors="replace")


def csv_table(data: bytes) -> list[list[str]]:
    """Wiersze CSV z rozpoznanym separatorem (przecinek, średnik, tabulator)."""
    text = decode_csv(data)
    sample = "\n".join(text.splitlines()[:40])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        counts = {sep: sample.count(sep) for sep in (",", ";", "\t")}
        delimiter = max(counts, key=counts.get) if any(counts.values()) else ","
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def xlsx_tables(data: bytes):
    """Karty XLSX jako tabele **surowych** wartości (czas trwania bywa ``timedelta`` albo ``time``).

    Pierwsza karta to spotkanie główne, kolejne – pokoje podgrup; generator oddaje je po kolei,
    a ``parse_bytes`` bierze pierwszą z rozpoznanym nagłówkiem. Strażnik bomby ZIP przed openpyxl
    z tego samego powodu, co w imporcie uczniów (``apps.core.xlsx``).
    """
    from openpyxl import load_workbook

    from apps.core.xlsx import XlsxBombError, check_xlsx_bomb

    try:
        check_xlsx_bomb(data)
    except XlsxBombError as exc:
        raise MeetReportError("Plik XLSX jest za duży po rozpakowaniu albo uszkodzony. Wgraj CSV.") from exc
    try:
        workbook = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl podnosi kilkanaście różnych klas
        raise MeetReportError(
            "Nie udało się otworzyć pliku XLSX. Zapisz go ponownie albo wgraj CSV."
        ) from exc
    try:
        for name in workbook.sheetnames:
            table = []
            for row in workbook[name].iter_rows(values_only=True):
                if not any(cell not in (None, "") for cell in row):
                    continue
                table.append(list(row))
                # Nagłówek, blok meta i limit wierszy z zapasem – reszty arkusza nie trzymamy.
                if len(table) > MAX_REPORT_ROWS + HEADER_SCAN_ROWS + 1:
                    break
            yield table
    finally:
        workbook.close()


def parse_bytes(data: bytes, filename: str = "") -> MeetReport:
    """Bajty pliku → raport. XLSX po sygnaturze ZIP albo rozszerzeniu, reszta jako CSV."""
    if len(data) > MAX_REPORT_BYTES:
        raise MeetReportError(f"Plik jest za duży (maks. {MAX_REPORT_BYTES // (1024 * 1024)} MB).")
    if not data.strip():
        raise MeetReportError("Plik jest pusty.")
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04" or name.endswith((".xlsx", ".xlsm")):
        first_error = None
        for table in xlsx_tables(data):
            try:
                return parse_table(table, filename)
            except MeetReportError as exc:
                first_error = first_error or exc
        raise first_error or MeetReportError("Plik XLSX nie ma żadnej karty z raportem.")
    if name.endswith((".xls", ".ods", ".pdf")):
        raise MeetReportError("Obsługujemy pliki CSV i XLSX. Zapisz raport w jednym z tych formatów.")
    return parse_table(csv_table(data), filename)


def read_upload(upload) -> MeetReport:
    """Wgrany plik → raport; rozmiar sprawdzany przed odczytem całości."""
    size = getattr(upload, "size", None)
    if size is not None and size > MAX_REPORT_BYTES:
        raise MeetReportError(f"Plik jest za duży (maks. {MAX_REPORT_BYTES // (1024 * 1024)} MB).")
    data = upload.read(MAX_REPORT_BYTES + 1)
    return parse_bytes(data, getattr(upload, "name", "") or "")


# --- link do arkusza -----------------------------------------------------------------------------


@dataclass(frozen=True)
class SheetLink:
    sheet_id: str
    gid: str

    @property
    def export_url(self) -> str:
        return f"https://{SHEETS_HOST}/spreadsheets/d/{self.sheet_id}/export?format=csv&gid={self.gid}"


def parse_sheet_link(url: str) -> SheetLink:
    """Wklejony adres arkusza → identyfikator i karta. Każdy inny adres to odmowa (WS-ATT-02 § 2.1).

    Serwer **nigdy** nie pobiera adresu wklejonego przez człowieka – z wklejonego bierze tylko ID
    i ``gid`` i składa adres eksportu sam. Sprawdzenie hosta jest więc drugą linią obrony, a nie
    jedyną: nawet przeoczony wariant zapisu nie zmieni hosta, z którym serwer się łączy.
    """
    text = (url or "").strip().strip("<>\"'")
    if not text:
        raise MeetReportError("Wklej link do arkusza Google.")
    try:
        parts = urlsplit(text)
        port = parts.port
    except ValueError as exc:
        raise MeetReportError("To nie jest poprawny adres.") from exc
    host = (parts.hostname or "").lower()
    if parts.scheme != "https" or host != SHEETS_HOST or port is not None or parts.username or parts.password:
        raise MeetReportError(
            "Akceptujemy wyłącznie link do arkusza Google w postaci https://docs.google.com/spreadsheets/d/…"
        )
    if PUBLISHED_PATH.match(parts.path):
        raise MeetReportError(
            "To jest link „opublikuj w internecie”. Wklej link z paska adresu arkusza "
            "(https://docs.google.com/spreadsheets/d/…/edit)."
        )
    match = SHEET_PATH.match(parts.path)
    if not match:
        raise MeetReportError(
            "Nie rozpoznano identyfikatora arkusza. Wklej link z paska adresu: "
            "https://docs.google.com/spreadsheets/d/…/edit"
        )
    gid = (parse_qs(parts.query).get("gid") or [""])[0]
    if not gid:
        gid = (parse_qs(parts.fragment).get("gid") or [""])[0]
    if gid and not GID.match(gid):
        raise MeetReportError("Numer karty arkusza (gid) w linku jest nieprawidłowy.")
    return SheetLink(sheet_id=match.group("id"), gid=gid or "0")


def _http_get(url: str):
    """Jedno żądanie GET bez podążania za przekierowaniem. Osobna funkcja – podmieniają ją testy."""
    import requests

    return requests.get(
        url,
        timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
        allow_redirects=False,
        stream=True,
        headers={"User-Agent": "olimpiada-workshop-attendance/1.0", "Accept": "text/csv, */*;q=0.1"},
    )


def _allowed_redirect(url: str) -> bool:
    parts = urlsplit(url)
    host = (parts.hostname or "").lower()
    if parts.scheme != "https" or parts.username or parts.password:
        return False
    try:
        if parts.port is not None:
            return False
    except ValueError:
        return False
    if host.endswith(EXPORT_HOST_SUFFIX):
        return True
    return host == SHEETS_HOST and parts.path.startswith("/spreadsheets/")


def _is_login(url: str) -> bool:
    return (urlsplit(url).hostname or "").lower() in LOGIN_HOSTS


def _read_body(response, deadline: float) -> bytes:
    """Treść odpowiedzi z limitem rozmiaru i czasu całości."""
    declared = response.headers.get("Content-Length")
    if declared and declared.isdigit() and int(declared) > MAX_REPORT_BYTES:
        raise MeetReportError(
            f"Arkusz jest za duży (ponad {MAX_REPORT_BYTES // (1024 * 1024)} MB). {DOWNLOAD_INSTEAD}"
        )
    chunks = []
    total = 0
    for chunk in response.iter_content(chunk_size=65536):
        if not chunk:
            continue
        total += len(chunk)
        if total > MAX_REPORT_BYTES:
            raise MeetReportError(
                f"Arkusz jest za duży (ponad {MAX_REPORT_BYTES // (1024 * 1024)} MB). {DOWNLOAD_INSTEAD}"
            )
        if time.monotonic() > deadline:
            raise MeetReportError(f"Pobieranie arkusza trwa zbyt długo. {DOWNLOAD_INSTEAD}")
        chunks.append(chunk)
    return b"".join(chunks)


def fetch_sheet_csv(link: SheetLink) -> bytes:
    """Eksport CSV karty arkusza udostępnionego „każdy, kto ma link” (WS-ATT-02 § 2.3–2.5).

    Przekierowania sprawdzamy ręcznie, bo ``requests`` poszedłby za każdym – a jedynymi celami,
    które ma prawo dostać ten eksport, są serwery Google oddające pliki. Strona logowania Google
    (przekierowanie albo odpowiedź HTML) znaczy, że arkusz nie jest udostępniony: to najczęstszy
    przypadek i dlatego ma własny, instruktażowy komunikat.
    """
    import requests

    url = link.export_url
    deadline = time.monotonic() + TOTAL_TIMEOUT
    for _ in range(MAX_REDIRECTS + 1):
        try:
            response = _http_get(url)
        except requests.Timeout as exc:
            raise MeetReportError(f"Google nie odpowiedział na czas. {DOWNLOAD_INSTEAD}") from exc
        except requests.RequestException as exc:
            raise MeetReportError(f"Nie udało się połączyć z Google. {DOWNLOAD_INSTEAD}") from exc
        try:
            status = response.status_code
            if status in (301, 302, 303, 307, 308):
                target = urljoin(url, response.headers.get("Location") or "")
                if _is_login(target):
                    raise MeetReportError(NOT_SHARED_MESSAGE)
                if not _allowed_redirect(target):
                    raise MeetReportError(
                        f"Google przekierował pobieranie pod nieoczekiwany adres. {DOWNLOAD_INSTEAD}"
                    )
                url = target
                continue
            if status in (401, 403):
                raise MeetReportError(NOT_SHARED_MESSAGE)
            if status == 404:
                raise MeetReportError("Nie ma takiego arkusza (albo karty o tym numerze). Sprawdź link.")
            if status != 200:
                raise MeetReportError(f"Google odpowiedział błędem {status}. {DOWNLOAD_INSTEAD}")
            content_type = (response.headers.get("Content-Type") or "").lower()
            if "text/html" in content_type:
                raise MeetReportError(NOT_SHARED_MESSAGE)
            try:
                body = _read_body(response, deadline)
            except requests.Timeout as exc:
                raise MeetReportError(f"Google nie odpowiedział na czas. {DOWNLOAD_INSTEAD}") from exc
            except requests.RequestException as exc:
                raise MeetReportError(f"Pobieranie arkusza przerwane. {DOWNLOAD_INSTEAD}") from exc
            # Strona logowania potrafi przyjść z 200 i bez nagłówka typu – poznajemy ją po treści.
            if body.lstrip()[:15].lower().startswith((b"<!doctype html", b"<html")):
                raise MeetReportError(NOT_SHARED_MESSAGE)
            return body
        finally:
            response.close()
    raise MeetReportError(f"Za dużo przekierowań przy pobieraniu arkusza. {DOWNLOAD_INSTEAD}")


def read_link(url: str) -> MeetReport:
    """Link do arkusza → raport (walidacja, pobranie, odczyt)."""
    link = parse_sheet_link(url)
    return parse_bytes(fetch_sheet_csv(link), "")
